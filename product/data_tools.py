"""Typed, bounded dataframe tools. Original inputs are never overwritten."""
from datetime import datetime
from decimal import Decimal, InvalidOperation, localcontext
import io
import json
from pathlib import Path
import re
import secrets
from typing import Literal
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from product.excel_report import put, numeric_decimal
from product.storage import digest, encode, timestamp

MAX_ROWS, MAX_CELLS, MAX_FRAMES = 50000, 250000, 60


class Args(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

class Empty(Args):
    pass

class Frame(Args):
    frame_id: str

class Load(Args):
    file_id: str
    sheet: str | int = 0
    header_row: int = Field(default=1, ge=1, le=20)

class Columns(Frame):
    columns: list[str] = Field(min_length=1, max_length=200)

class Limit(Frame):
    limit: int = Field(ge=1, le=5000)
    offset: int = Field(default=0, ge=0, le=50000)

class Rename(Frame):
    mapping: dict[str, str] = Field(min_length=1, max_length=200)

class Condition(Args):
    column: str
    op: Literal["eq", "ne", "gt", "gte", "lt", "lte", "contains", "starts_with", "in", "is_empty", "not_empty"]
    value: str | int | float | list[str] | None = None
    kind: Literal["text", "number", "date"] = "text"

class Filter(Frame):
    conditions: list[Condition] = Field(min_length=1, max_length=10)
    combine: Literal["and", "or"] = "and"

class Sort(Frame):
    by: list[str] = Field(min_length=1, max_length=5)
    descending: bool = False
    kinds: dict[str, Literal["text", "number", "date"]] = Field(default_factory=dict)

class Dedupe(Frame):
    keys: list[str] = Field(min_length=1, max_length=20)
    keep: Literal["first", "last"] = "first"
    duplicates_only: bool = False

class Fill(Frame):
    values: dict[str, str] = Field(min_length=1, max_length=30)

class Normalize(Columns):
    mode: Literal["trim", "lower", "upper"] = "trim"

class Concat(Args):
    frame_ids: list[str] = Field(min_length=2, max_length=20)
    union_columns: bool = False

class Join(Args):
    left_id: str
    right_id: str
    left_on: list[str] = Field(min_length=1, max_length=5)
    right_on: list[str] = Field(min_length=1, max_length=5)
    how: Literal["inner", "left", "outer"] = "left"
    relationship: Literal["one_to_one", "many_to_one"] = "one_to_one"

class Metric(Args):
    column: str
    op: Literal["sum", "mean", "min", "max", "count", "nunique"]
    output: str

class Aggregate(Frame):
    group_by: list[str] = Field(default_factory=list, max_length=4)
    metrics: list[Metric] = Field(min_length=1, max_length=10)

class Derive(Frame):
    column: str
    operation: Literal["add", "subtract", "multiply", "divide", "month", "year", "absolute"]
    left: str
    right: str | None = None
    factor: str | int | float | None = None

class Chart(Frame):
    x: str
    y: str
    kind: Literal["bar", "line", "pie"] = "bar"
    title: str = Field(default="Analysis", max_length=120)

class Export(Args):
    frame_ids: list[str] = Field(min_length=1, max_length=8)
    title: str = Field(default="Data analysis", min_length=1, max_length=120)
    include_charts: bool = True


SPECS = {
    "list_files": (Empty, "List captured CSV/XLSX files and their file_id."),
    "inspect_table": (Load, "Read file/sheet headers, row count and 5 sample rows; returns a reusable frame_id."),
    "profile_table": (Frame, "Count missing values, duplicate rows and numeric column ranges."),
    "select_columns": (Columns, "Keep the named columns."),
    "limit_rows": (Limit, "Return up to limit rows after offset. For top N, sort first and then limit_rows."),
    "rename_columns": (Rename, "Rename columns using an explicit old-to-new mapping; result names must stay unique."),
    "filter_rows": (Filter, "Filter with conditions {column,op,value,kind}; combine=and/or; date values must be ISO."),
    "sort_rows": (Sort, "Sort by columns; kinds maps numeric/date columns to number/date; descending=true for highest first."),
    "deduplicate": (Dedupe, "Remove duplicates by explicit keys; duplicates_only=true shows all duplicate occurrences instead."),
    "fill_missing": (Fill, "Fill blank cells using an explicit column-to-text values mapping."),
    "normalize_text": (Normalize, "Trim whitespace or change letter case in selected columns."),
    "concatenate": (Concat, "Append frame_ids; union_columns=true explicitly allows different column sets."),
    "join_tables": (Join, "Join left_id/right_id on explicit keys; validate one_to_one or many_to_one."),
    "aggregate": (Aggregate, "Group by columns and compute metrics {column,op,output}; op=sum/mean/min/max/count/nunique."),
    "derive_column": (Derive, "Create a NEW column: operation add/subtract/multiply/divide with left and right column OR factor; month/year/absolute use left only."),
    "create_chart": (Chart, "Plot frame_id with x/y columns; kind=bar/line/pie; maximum 30 points. Aggregate first."),
    "export_report": (Export, "Generate an Excel report of frame_ids, optionally including charts. Returns a downloadable artifact."),
}


def number(value):
    text = str(value).strip()
    # Reject malformed grouping, currency symbols and non-finite numbers.
    if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", text):
        raise ValueError(f"Not an unambiguous number: {text[:60]!r}. Clean the column or choose a different field.")
    parsed = Decimal(text.replace(",", ""))
    if len(parsed.as_tuple().digits) > 60 or abs(parsed) > Decimal("1e30"):
        raise ValueError("A numeric value exceeds the supported precision or magnitude.")
    return parsed


def day(value):
    text = str(value).strip().replace("/", "-")
    if not re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])", text):
        raise ValueError("Dates must use YYYY-MM-DD; ambiguous dates require clarification.")
    return datetime.fromisoformat(text).date()


def converted(value, kind):
    return number(value) if kind == "number" else day(value) if kind == "date" else str(value)


class DataTools:
    def __init__(self, assistant, chat_id, turn_id, state):
        self.assistant, self.service = assistant, assistant.service
        self.store, self.chat_id, self.turn_id, self.state = assistant.store, chat_id, turn_id, state

    def call(self, tool, arguments):
        if tool not in SPECS:
            raise ValueError("Unknown tool. Choose a supported data tool.")
        try: args = SPECS[tool][0].model_validate(arguments)
        except ValidationError as exc:
            problems = "; ".join(".".join(str(part) for part in issue["loc"]) + ": " + ("unexpected argument" if issue["type"] == "extra_forbidden" else issue["msg"]) for issue in exc.errors())
            raise ValueError(f"Invalid arguments for {tool}: {problems}. Allowed argument names: {', '.join(SPECS[tool][0].model_fields)}. Retry with only those arguments.") from exc
        with localcontext() as context:
            context.prec = 80
            return getattr(self, tool)(**args.model_dump())

    def frame(self, frame_id):
        info = self.state["frames"].get(frame_id)
        if not info:
            raise ValueError(f"Unknown frame_id {frame_id!r}. Inspect a file first or use an existing frame ID.")
        path = self.service.path(info["path"])
        if digest(path) != info["sha256"]:
            raise ValueError("A saved result changed; its integrity check failed.")
        data = json.loads(path.read_text(encoding="utf-8"))
        return pd.DataFrame(data["rows"], columns=data["columns"], dtype=object)

    def columns(self, frame, columns):
        missing = set(columns) - set(frame.columns)
        if missing:
            raise ValueError("Unknown columns: " + ", ".join(sorted(missing)))

    def save(self, frame, operation, parents=()):
        if len(frame) > MAX_ROWS or frame.size > MAX_CELLS or len(frame.columns) > 200:
            raise ValueError("Result exceeds 50,000 rows, 200 columns or 250,000 cells. Select fewer rows/columns.")
        if len(self.state["frames"]) >= MAX_FRAMES:
            raise ValueError("This conversation has reached 60 results. Start a new conversation.")
        if sum(info["rows"] for info in self.state["frames"].values()) + len(frame) > 200000:
            raise ValueError("This conversation has reached its result-row limit. Start a new conversation.")
        frame = frame.fillna("").map(lambda value: str(value))
        frame_id = f"t{len(self.state['frames']) + 1}"
        relative = f"files/assistant/{self.chat_id}/frames/{frame_id}.json"
        path = self.service.path(relative)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(encode({"columns": list(frame.columns), "rows": frame.values.tolist()}), encoding="utf-8")
        info = {"id": frame_id, "operation": operation, "parents": list(parents), "rows": len(frame),
                "columns": list(frame.columns), "path": relative, "sha256": digest(path)}
        currencies = {str(value).strip().upper() for column in frame.columns if re.sub(r"[^a-z]", "", str(column).lower()) in ("currency", "currencycode", "ccy") for value in frame[column] if str(value).strip()}
        has_currency_column = any(re.sub(r"[^a-z]", "", str(column).lower()) in ("currency", "currencycode", "ccy") for column in frame.columns)
        if not has_currency_column: currencies.update(currency for parent in parents for currency in self.state["frames"].get(parent, {}).get("currencies", []))
        info["currencies"] = sorted(currencies)
        info["difference_columns"] = [column for column in frame.columns if any(column in self.state["frames"].get(parent, {}).get("difference_columns", []) for parent in parents)]
        info["group"] = self.state["frames"].get(parents[0], {}).get("group", "input") if parents else "input"
        info["source_groups"] = sorted({group for parent in parents for group in self.state["frames"].get(parent, {}).get("source_groups", ["input"])}) if parents else ["input"]
        self.state["frames"][frame_id] = info
        self.state["latest_frame"] = frame_id
        return self.preview(frame_id, frame)

    def preview(self, frame_id, frame=None):
        frame = self.frame(frame_id) if frame is None else frame
        return {"frame_id": frame_id, "rows": len(frame), "columns": list(frame.columns),
                "preview": frame.head(5).values.tolist()}

    def list_files(self):
        return {"files": [{**{key: file[key] for key in ("file_id", "name", "sheets")}, "group": file.get("group", "input")} for file in self.assistant.catalog(self.chat_id)]}

    def inspect_table(self, file_id, sheet=0, header_row=1):
        item = next((item for item in self.assistant.catalog(self.chat_id) if item["file_id"] == file_id), None)
        if not item:
            raise ValueError("Unknown file_id. Use list_files first.")
        cache = f"{file_id}:{sheet}:{header_row}"
        for frame_id, info in self.state["frames"].items():
            if info.get("source") == cache:
                return self.preview(frame_id)
        path = self.service.path(item["path"])
        if digest(path) != item["sha256"]:
            raise ValueError("The captured input failed its integrity check.")
        if item["name"].lower().endswith(".csv"):
            if sheet not in (0, "CSV"):
                raise ValueError("CSV files have one table named CSV.")
            frame = pd.read_csv(path, dtype=str, keep_default_na=False, header=None, nrows=5021, encoding="utf-8-sig")
        else:
            if isinstance(sheet, int) and not 0 <= sheet < len(item["sheets"]) or isinstance(sheet, str) and sheet not in item["sheets"]:
                raise ValueError("Unknown worksheet. Use a name returned by list_files.")
            frame = pd.read_excel(path, sheet_name=sheet, dtype=str, keep_default_na=False, header=None, nrows=5021, engine="openpyxl")
        if len(frame) < header_row:
            raise ValueError("The requested header row does not exist.")
        columns = [str(value).strip() for value in frame.iloc[header_row - 1]]
        if not columns or any(not column for column in columns) or len(set(columns)) != len(columns):
            raise ValueError("Headers are blank or duplicated. Choose the correct header_row or fix the source copy.")
        frame = frame.iloc[header_row:].copy()
        frame.columns = columns
        if len(frame) > 5000:
            raise ValueError("One input table exceeds 5,000 rows. Split the input before importing.")
        result = self.save(frame.reset_index(drop=True), "Inspect " + item["name"])
        self.state["frames"][result["frame_id"]].update(source=cache, group=item.get("group", "input"), source_groups=[item.get("group", "input")])
        return result

    def profile_table(self, frame_id):
        frame = self.frame(frame_id)
        columns = []
        for column in frame.columns:
            values = frame[column].astype(str)
            present = values[values.str.strip() != ""]
            info = {"column": column, "missing": len(values) - len(present), "distinct": int(present.nunique())}
            if len(present):
                try:
                    numbers = [number(value) for value in present]
                    info.update(min=str(min(numbers)), max=str(max(numbers)))
                except ValueError:
                    pass
            columns.append(info)
        return {"frame_id": frame_id, "rows": len(frame), "duplicate_rows": int(frame.duplicated().sum()), "column_profile": columns}

    def select_columns(self, frame_id, columns):
        frame = self.frame(frame_id)
        self.columns(frame, columns)
        if len(set(columns)) != len(columns):
            raise ValueError("Choose each column once.")
        return self.save(frame[columns], "Select columns", [frame_id])

    def filter_rows(self, frame_id, conditions, combine="and"):
        frame = self.frame(frame_id)
        self.columns(frame, [condition["column"] for condition in conditions])
        mask = pd.Series(combine == "and", index=frame.index)
        for condition in conditions:
            series = frame[condition["column"]].astype(str)
            op, value, kind = condition["op"], condition["value"], condition["kind"]
            if op in ("is_empty", "not_empty"):
                test = series.str.strip() == ""
                if op == "not_empty": test = ~test
            elif op == "in":
                if not isinstance(value, list): raise ValueError("The in filter requires a list of text values.")
                test = series.isin(value)
            elif op in ("contains", "starts_with"):
                if not isinstance(value, str): raise ValueError("Text filters require a text value.")
                test = series.str.contains(value, regex=False) if op == "contains" else series.str.startswith(value)
            else:
                if value is None or isinstance(value, list): raise ValueError("Comparison requires a single value.")
                target = converted(value, kind)
                def compare(raw):
                    if not str(raw).strip() and kind != "text": return False
                    item = converted(raw, kind)
                    return {"eq": lambda: item == target, "ne": lambda: item != target,
                            "gt": lambda: item > target, "gte": lambda: item >= target,
                            "lt": lambda: item < target, "lte": lambda: item <= target}[op]()
                test = series.map(compare)
            mask = mask & test if combine == "and" else mask | test
        return self.save(frame[mask].reset_index(drop=True), "Filter rows", [frame_id])

    def limit_rows(self, frame_id, limit, offset=0):
        return self.save(self.frame(frame_id).iloc[offset:offset + limit].reset_index(drop=True), "Select row range", [frame_id])

    def rename_columns(self, frame_id, mapping):
        frame = self.frame(frame_id); self.columns(frame, mapping)
        names = [mapping.get(column, column) for column in frame.columns]
        if any(not name.strip() for name in names) or len(set(names)) != len(names): raise ValueError("New column names must be nonempty and unique.")
        return self.save(frame.rename(columns=mapping), "Rename columns", [frame_id])

    def sort_rows(self, frame_id, by, descending=False, kinds=None):
        frame = self.frame(frame_id)
        self.columns(frame, by)
        if set(kinds or {}) - set(by): raise ValueError("Sort kinds must refer to sort columns.")
        # Stable sorting keeps empty cells last, including descending numeric sorts.
        indexes = list(frame.index)
        for column in reversed(by):
            present = [i for i in indexes if str(frame.at[i, column]).strip()]
            missing = [i for i in indexes if not str(frame.at[i, column]).strip()]
            indexes = sorted(present, key=lambda i: converted(frame.at[i, column], (kinds or {}).get(column, "text")), reverse=descending) + missing
        return self.save(frame.loc[indexes].reset_index(drop=True), "Sort rows", [frame_id])

    def deduplicate(self, frame_id, keys, keep="first", duplicates_only=False):
        frame = self.frame(frame_id)
        self.columns(frame, keys)
        result = frame[frame.duplicated(keys, keep=False)] if duplicates_only else frame.drop_duplicates(keys, keep=keep)
        return self.save(result.reset_index(drop=True), "Find duplicates" if duplicates_only else "Remove duplicates", [frame_id])

    def fill_missing(self, frame_id, values):
        frame = self.frame(frame_id).copy()
        self.columns(frame, values)
        for column, value in values.items():
            frame.loc[frame[column].astype(str).str.strip() == "", column] = value
        return self.save(frame, "Fill missing values", [frame_id])

    def normalize_text(self, frame_id, columns, mode="trim"):
        frame = self.frame(frame_id).copy()
        self.columns(frame, columns)
        for column in columns:
            series = frame[column].astype(str)
            frame[column] = series.str.strip() if mode == "trim" else series.str.lower() if mode == "lower" else series.str.upper()
        return self.save(frame, "Normalize text", [frame_id])

    def concatenate(self, frame_ids, union_columns=False):
        if len(set(frame_ids)) != len(frame_ids): raise ValueError("Appending a table to itself requires a different explicit workflow.")
        groups = {self.state["frames"].get(frame_id, {}).get("group", "input") for frame_id in frame_ids}
        if len(groups) > 1: raise ValueError("Input and BI reference tables must not be appended together. Use join_tables with explicit keys to compare them.")
        frames = [self.frame(frame_id) for frame_id in frame_ids]
        if sum(len(frame) for frame in frames) > MAX_ROWS: raise ValueError("The merged result would exceed 50,000 rows.")
        if not union_columns and any(set(frame.columns) != set(frames[0].columns) for frame in frames):
            raise ValueError("Column sets differ. Explicitly allow union_columns or select matching columns first.")
        columns = set().union(*(set(frame.columns) for frame in frames))
        if len(columns) * sum(len(frame) for frame in frames) > MAX_CELLS: raise ValueError("The merged result would exceed the cell limit.")
        return self.save(pd.concat(frames, ignore_index=True, sort=False).fillna(""), "Append tables", frame_ids)

    def join_tables(self, left_id, right_id, left_on, right_on, how="left", relationship="one_to_one"):
        left, right = self.frame(left_id), self.frame(right_id)
        self.columns(left, left_on); self.columns(right, right_on)
        if len(left_on) != len(right_on): raise ValueError("Join key lists must have the same length.")
        if any(frame[keys].map(lambda value: not str(value).strip()).any().any() for frame, keys in ((left, left_on), (right, right_on))):
            raise ValueError("Join keys contain missing values. Clean or exclude those rows first.")
        if right.duplicated(right_on).any() or relationship == "one_to_one" and left.duplicated(left_on).any():
            raise ValueError("Duplicate join keys violate the requested relationship. Resolve duplicates before joining.")
        if len(left) + (len(right) if how == "outer" else 0) > MAX_ROWS:
            raise ValueError("The projected join exceeds the row limit.")
        result = left.merge(right, left_on=left_on, right_on=right_on, how=how, validate=relationship, suffixes=("_left", "_right"))
        if len(set(result.columns)) != len(result.columns): raise ValueError("Join produced duplicate column names. Select or rename source fields first.")
        return self.save(result.fillna(""), "Join tables", [left_id, right_id])

    def aggregate(self, frame_id, group_by, metrics):
        frame = self.frame(frame_id)
        self.columns(frame, group_by + [metric["column"] for metric in metrics])
        outputs = [metric["output"] for metric in metrics]
        if len(set(group_by + outputs)) != len(group_by + outputs) or any(not output.strip() for output in outputs):
            raise ValueError("Group and output column names must be unique and nonempty.")
        currency_columns = [column for column in frame.columns if re.sub(r"[^a-z]", "", column.lower()) in ("currency", "currencycode", "ccy")]
        if not currency_columns and len(self.state["frames"][frame_id].get("currencies", [])) > 1 and any(metric["op"] in ("sum", "mean", "min", "max") for metric in metrics):
            raise ValueError("Multiple currencies remain in this result's source evidence. Return to a table with a currency column and group by currency.")
        for column in frame.columns:
            if re.sub(r"[^a-z]", "", column.lower()) in ("currency", "currencycode", "ccy"):
                if frame[column].nunique() > 1 and column not in group_by and any(m["op"] in ("sum", "mean", "min", "max") for m in metrics):
                    raise ValueError("Multiple currencies are present. Group by the currency column before numeric aggregation.")
        groups = frame.groupby(group_by, dropna=False, sort=False) if group_by else [((), frame)]
        rows = []
        for key, part in groups:
            keys = list(key if isinstance(key, tuple) else (key,)) if group_by else []
            values = []
            for metric in metrics:
                items = [value for value in part[metric["column"]] if str(value).strip()]
                op = metric["op"]
                if op == "count": value = len(items)
                elif op == "nunique": value = len(set(items))
                elif not items: value = ""
                else:
                    numbers = [number(item) for item in items]
                    value = sum(numbers, Decimal(0)) if op == "sum" else sum(numbers, Decimal(0)) / len(numbers) if op == "mean" else min(numbers) if op == "min" else max(numbers)
                values.append(str(value))
            rows.append(keys + values)
        result = self.save(pd.DataFrame(rows, columns=group_by + outputs), "Aggregate", [frame_id])
        differences = self.state["frames"][frame_id].get("difference_columns", [])
        self.state["frames"][result["frame_id"]]["difference_columns"] = [metric["output"] for metric in metrics if metric["column"] in differences and metric["op"] in ("sum", "mean", "min", "max")]
        return result

    def derive_column(self, frame_id, column, operation, left, right=None, factor=None):
        frame = self.frame(frame_id).copy()
        self.columns(frame, [left] + ([right] if right else []))
        if column in frame or not column.strip(): raise ValueError("Choose a new, nonempty column name.")
        if operation in ("add", "subtract", "multiply", "divide") and (right is None) == (factor is None):
            raise ValueError("Specify either a right-hand column or a numeric factor.")
        def derive(row):
            if operation in ("year", "month"):
                date = day(row[left]); return str(date.year) if operation == "year" else date.strftime("%Y-%m")
            a = number(row[left])
            if operation == "absolute": return str(abs(a))
            b = number(row[right]) if right else number(factor)
            if operation == "divide" and not b: raise ValueError("Division by zero. Clean or filter these rows first.")
            return str(a + b if operation == "add" else a - b if operation == "subtract" else a * b if operation == "multiply" else a / b)
        frame[column] = frame.apply(derive, axis=1) if len(frame) else pd.Series(dtype=object)
        result = self.save(frame, "Derive " + column, [frame_id])
        if operation == "subtract": self.state["frames"][result["frame_id"]]["difference_columns"].append(column)
        return result

    def create_chart(self, frame_id, x, y, kind="bar", title="Analysis"):
        frame = self.frame(frame_id); self.columns(frame, [x, y])
        if not 1 <= len(frame) <= 30: raise ValueError("Charts require 1 to 30 rows. Aggregate or filter the table first.")
        if frame[x].duplicated().any(): raise ValueError("Chart labels repeat. Aggregate to one row per label first.")
        values = [number(value) for value in frame[y]]
        if any(abs(value) > Decimal("1e13") for value in values): raise ValueError("Values are too large for a reliable chart.")
        if kind == "pie" and (any(value < 0 for value in values) or not sum(values)):
            raise ValueError("Pie charts require nonnegative values with a positive total.")
        if len(self.state["charts"]) >= 20: raise ValueError("This conversation already contains 20 charts.")
        chart_id = f"c{len(self.state['charts']) + 1}"
        chart = {"chart_id": chart_id, "frame_id": frame_id, "x": x, "y": y, "kind": kind, "title": title,
                 "labels": frame[x].tolist(), "values": [str(value) for value in values]}
        self.state["charts"][chart_id] = chart
        return chart

    def export_report(self, frame_ids, title="Data analysis", include_charts=True):
        frames = [self.frame(frame_id) for frame_id in frame_ids]
        if sum(frame.size for frame in frames) > MAX_CELLS: raise ValueError("The report exceeds 250,000 cells. Export a smaller selection.")
        workbook = Workbook(); summary = workbook.active; summary.title = "Summary"
        put(summary, 1, 1, title).font = Font(name="Calibri", size=20, bold=True, color="292B2E")
        summary.column_dimensions["A"].width = 24; summary.column_dimensions["B"].width = 100
        turn = self.store.one("SELECT instruction FROM data_turns WHERE id=?", (self.turn_id,))
        for row, (label, value) in enumerate([( "Request", turn["instruction"]), ("Generated UTC", timestamp()),
                                             ("Scope", "Local captured files; derived results do not change original inputs.")], 3):
            put(summary, row, 1, label); put(summary, row, 2, value); summary.row_dimensions[row].height = 45
        chat = self.assistant.chat(self.chat_id)
        for row, (label, value) in enumerate([("Input folder", chat["folder"]), ("BI system folder", chat["bi_folder"] or "Not selected")], 6):
            put(summary, row, 1, label); put(summary, row, 2, value)
            summary.row_dimensions[row].height = 35
        for index, (frame_id, frame) in enumerate(zip(frame_ids, frames), 1):
            sheet = workbook.create_sheet(f"Result {index}"); sheet.freeze_panes = "A2"
            for col, column in enumerate(frame.columns, 1):
                put(sheet, 1, col, column).font = Font(name="Calibri", bold=True, color="FFFFFF")
                sheet.cell(1, col).fill = PatternFill("solid", fgColor="455B72")
                sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 22
            # All source/derived values remain exact text, preserving identifiers and decimal precision.
            for rownum, row in enumerate(frame.values.tolist(), 2):
                for col, value in enumerate(row, 1): put(sheet, rownum, col, value)
            if len(frame.columns): sheet.auto_filter.ref = sheet.dimensions
            put(summary, 7 + index, 1, frame_id); put(summary, 7 + index, 2, f"Result {index}: {len(frame)} rows; " + self.state["frames"][frame_id]["operation"])
        if include_charts:
            for index, chart in enumerate(self.state["charts"].values(), 1):
                if chart["frame_id"] not in frame_ids: continue
                sheet = workbook.create_sheet(f"Chart {index}")
                put(sheet, 1, 1, chart["x"]); put(sheet, 1, 2, chart["y"])
                for row, (label, value) in enumerate(zip(chart["labels"], chart["values"]), 2):
                    put(sheet, row, 1, label); put(sheet, row, 2, float(number(value)))
                graph = BarChart() if chart["kind"] == "bar" else LineChart() if chart["kind"] == "line" else PieChart()
                graph.title = chart["title"]; graph.style = 13
                graph.add_data(Reference(sheet, min_col=2, min_row=1, max_row=1 + len(chart["values"])), titles_from_data=True)
                graph.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=1 + len(chart["values"])))
                sheet.add_chart(graph, "D2")
        audit = workbook.create_sheet("Source evidence")
        for column, title_text in enumerate(["Group", "File", "SHA-256"], 1): put(audit, 1, column, title_text)
        for row, file in enumerate(self.assistant.catalog(self.chat_id), 2):
            put(audit, row, 1, "BI system" if file.get("group") == "bi" else "Input"); put(audit, row, 2, file["name"]); put(audit, row, 3, file["sha256"])
        audit.column_dimensions["A"].width = 18; audit.column_dimensions["B"].width = 45; audit.column_dimensions["C"].width = 70
        audit.freeze_panes = "A2"
        operations = workbook.create_sheet("Operations")
        for col, label in enumerate(["Request", "Tool", "Arguments", "Status"], 1): put(operations, 1, col, label)
        rownum = 2
        for row in self.store.all("SELECT instruction,state_json FROM data_turns WHERE chat_id=? ORDER BY id", (self.chat_id,)):
            for step in json.loads(row["state_json"]).get("steps", []):
                for col, value in enumerate([row["instruction"], step["tool"], encode(step["arguments"]), step["status"]], 1): put(operations, rownum, col, value)
                rownum += 1
        for column in ("A", "B", "C", "D"): operations.column_dimensions[column].width = 40
        operations.freeze_panes = "A2"
        artifact_id = secrets.token_hex(12)
        relative = f"exports/Assistant-{artifact_id}.xlsx"; path = self.service.path(relative)
        workbook.save(path); workbook.close()
        artifact = {"id": artifact_id, "title": title, "name": path.name, "url": f"/api/assistant/artifacts/{artifact_id}", "frame_ids": frame_ids,
                    "row_count": sum(len(frame) for frame in frames), "turn_id": self.turn_id,
                    "chart_ids": [chart["chart_id"] for chart in self.state["charts"].values() if include_charts and chart["frame_id"] in frame_ids]}
        with self.store.connect() as db:
            db.execute("INSERT INTO data_artifacts VALUES(?,?,?,?,?,?,?)", (artifact_id, self.chat_id, self.turn_id, relative, digest(path), encode(artifact), timestamp()))
        return artifact
