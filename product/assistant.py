"""Persistent local data conversations and a bounded observe/act loop."""
import json
import os
from pathlib import Path
import re
import secrets
from typing import Literal
import zipfile

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from product.data_tools import DataTools, SPECS
from product.data_contracts import wants_difference_chart
from product.model import ModelUnavailable
from product.review_workflow import review_requested, run_review, review_summary, review_message
from product.input_folders import select_inputs, input_files
from product.storage import digest, encode, timestamp


class Action(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    action: Literal["tool", "final", "ask"]
    tool: str | None = None
    arguments: dict = Field(default_factory=dict)
    message: str = Field(default="", max_length=1600)


class Plan(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    supported: bool
    required_tools: list[str] = Field(max_length=10)
    message: str = Field(max_length=900)


def explicit_row_limit(instruction):
    """Check simple, measurable row bounds independently of model narration."""
    patterns = [r"\b(?:top|first)\s+(\d+)\s+(?:rows?|records?|orders?|suppliers?)\b",
                r"\b(?:keep|show|return|retain)\s+(?:only\s+)?(?:the\s+)?first\s+(row|record)\b",
                r"只(?:保留|取|要)(?:第)?([一二三四五六七八九十\d]+)(?:行|条)",
                r"前([一二三四五六七八九十\d]+)(?:行|条|名|个)"]
    digits = {letter: index for index, letter in enumerate("零一二三四五六七八九")}
    for pattern in patterns:
        match = re.search(pattern, instruction, re.I)
        if not match: continue
        before = instruction[max(0, match.start() - 12):match.start()].lower()
        if any(negative in before for negative in ("不要", "不需要", "别", "do not", "don't")): continue
        value = match.group(1)
        if value.isdigit(): return int(value)
        if value.lower() in ("row", "record"): return 1
        if "十" in value:
            left, right = value.split("十", 1)
            return (digits.get(left, 1) * 10) + digits.get(right, 0)
        return digits.get(value)
    return None


class DataAssistant:
    def __init__(self, service):
        self.service, self.store = service, service.store
        with self.store.connect() as db:
            for row in db.execute("SELECT id,state_json FROM data_turns WHERE status='running'").fetchall():
                state = json.loads(row[1]); state["error"] = "The application stopped during this request. Existing results are retained. Continue with a new message."
                db.execute("UPDATE data_turns SET status='failed',state_json=?,updated_at=? WHERE id=?", (encode(state), timestamp(), row[0]))

    def chat(self, chat_id):
        chat = self.store.one("SELECT * FROM data_chats WHERE id=?", (chat_id,))
        if not chat: raise HTTPException(404, "Conversation was not found.")
        return chat

    def catalog(self, chat_id):
        return json.loads(self.chat(chat_id)["catalog_json"])

    def list(self):
        return self.store.all("SELECT id,title,folder,bi_folder,created_at,updated_at FROM data_chats ORDER BY id DESC LIMIT 50")

    def create(self, folder, actor, *, bi_folder=None):
        selected, bi = select_inputs(folder, bi_folder, self.store.root)
        roots = [("input", selected)] + ([("bi", bi)] if bi else [])
        with self.service.guard:
            self.service.batches.idle()
            files, total = [], 0
            for group, origin_root in roots:
                group_count = 0
                for path in input_files(origin_root, self.store.root):
                    size = path.stat().st_size; total += size
                    if size > 10 * 1024 * 1024 or total > 100 * 1024 * 1024 or len(files) >= 100:
                        raise ValueError("Combined folder limit: 100 files, 10 MB per file and 100 MB in total.")
                    files.append((path, path.relative_to(origin_root).as_posix(), group, origin_root))
                    group_count += 1
                if not group_count: raise ValueError("No CSV or Excel files were found in the " + ("BI system folder." if group == "bi" else "input folder."))
            if not files: raise ValueError("No CSV or Excel files were found in this folder.")
            with self.store.connect() as db:
                cursor = db.execute("INSERT INTO data_chats(title,folder,catalog_json,state_json,actor,created_at,updated_at,bi_folder) VALUES(?,?,?,?,?,?,?,?)",
                                    (selected.name or "Data conversation", str(selected), "[]", encode({"frames": {}, "charts": {}, "latest_frame": None}), actor, timestamp(), timestamp(), str(bi) if bi else None))
                chat_id = cursor.lastrowid
            catalog, directory = [], self.service.path(f"files/assistant/{chat_id}")
            captured_total = 0
            directory.mkdir(parents=True, exist_ok=True)
            try:
                for index, (path, name, group, origin_root) in enumerate(files, 1):
                    if path.is_symlink() or not path.resolve().is_relative_to(origin_root): raise ValueError("An input path changed during capture.")
                    content = path.read_bytes()
                    captured_total += len(content)
                    if captured_total > 100 * 1024 * 1024: raise ValueError("Captured inputs exceed 100 MB.")
                    if len(content) > 10 * 1024 * 1024: raise ValueError("An input grew beyond 10 MB during capture.")
                    file_id = f"f{index}"; snapshot = directory / (file_id + path.suffix.lower()); snapshot.write_bytes(content)
                    if path.suffix.lower() == ".xlsx":
                        with zipfile.ZipFile(snapshot) as archive:
                            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                                raise ValueError("An Excel workbook expands beyond 100 MB.")
                        book = load_workbook(snapshot, read_only=True, data_only=True, keep_links=False)
                        try:
                            if len(book.sheetnames) > 20 or any(sheet.max_column > 200 for sheet in book):
                                raise ValueError("Excel limits: 20 worksheets and 200 columns.")
                            sheets = book.sheetnames
                        finally: book.close()
                    else: sheets = ["CSV"]
                    catalog.append({"file_id": file_id, "name": name, "sheets": sheets, "group": group,
                                    "path": snapshot.relative_to(self.store.root).as_posix(), "sha256": digest(snapshot)})
                with self.store.connect() as db:
                    db.execute("UPDATE data_chats SET catalog_json=? WHERE id=?", (encode(catalog), chat_id))
            except Exception:
                # Failed intake never leaves a usable partial conversation.
                with self.store.connect() as db: db.execute("DELETE FROM data_chats WHERE id=?", (chat_id,))
                for snapshot in directory.iterdir():
                    if snapshot.is_file(): snapshot.unlink()
                directory.rmdir()
                raise
        return self.describe(chat_id)

    def describe(self, chat_id):
        chat = self.chat(chat_id); state = json.loads(chat["state_json"])
        turns = []
        for row in self.store.all("SELECT * FROM data_turns WHERE chat_id=? ORDER BY id", (chat_id,)):
            turn = json.loads(row["state_json"])
            item = {"id": row["id"], "instruction": row["instruction"], "status": row["status"], **turn}
            item["can_retry"] = row["status"] in ("failed", "cancelled") or (row["status"] == "needs_input" and bool(chat["bi_folder"]) and review_requested(row["instruction"]))
            if turn.get("workflow") == "reconciliation_review" and turn.get("batch_id"):
                batch = self.service.batches.describe(turn["batch_id"])
                item["review_batch"] = review_summary(batch)
                if row["status"] != "running":
                    item["message"] = review_message(batch)
                    item["status"] = batch["status"]
                if batch["status"] == "cancelled": item["can_retry"] = False
            turns.append(item)
        preview = None
        if state["latest_frame"]:
            preview = DataTools(self, chat_id, None, state).preview(state["latest_frame"])
            preview["columns"] = preview["columns"][:30]
            preview["preview"] = [row[:30] for row in preview["preview"]]
        artifacts = [json.loads(row["info_json"]) for row in self.store.all("SELECT info_json FROM data_artifacts WHERE chat_id=? ORDER BY created_at", (chat_id,))]
        return {"id": chat_id, "title": chat["title"], "folder": chat["folder"], "bi_folder": chat["bi_folder"], "turns": turns,
                "files": [{**{key: item[key] for key in ("file_id", "name", "sheets")}, "group": item.get("group", "input")} for item in json.loads(chat["catalog_json"])],
                "frames": [{key: info[key] for key in ("id", "operation", "columns", "rows")} for info in state["frames"].values()],
                "preview": preview, "charts": list(state["charts"].values())[-5:], "artifacts": artifacts,
                "active_turn": self.service.assistant_active if any(turn["id"] == self.service.assistant_active for turn in turns) else None}

    def start(self, chat_id, instruction, actor, *, retry_turn_id=None):
        if not 3 <= len(instruction.strip()) <= 3000: raise ValueError("Enter a request between 3 and 3,000 characters.")
        with self.service.guard:
            self.service.batches.idle(); chat = self.chat(chat_id)
            if chat["actor"] != actor: raise HTTPException(403, "This conversation belongs to another account.")
            if len(self.store.all("SELECT id FROM data_turns WHERE chat_id=?", (chat_id,))) >= 50:
                raise ValueError("Start a new conversation after 50 requests.")
            initial = {"steps": [], "message": "", "error": None}
            if retry_turn_id is not None:
                previous = self.store.one("SELECT * FROM data_turns WHERE id=? AND chat_id=?", (retry_turn_id, chat_id))
                if not previous: raise HTTPException(404, "The request to retry was not found in this conversation.")
                if previous["status"] not in ("failed", "cancelled") and not (previous["status"] == "needs_input" and chat["bi_folder"] and review_requested(previous["instruction"])):
                    raise HTTPException(409, "Only a failed, stopped or previously rejected reconciliation request can be retried.")
                if instruction != previous["instruction"]: raise ValueError("Retry must keep the original request. Send a new message to change it.")
                prior_state = json.loads(previous["state_json"])
                if prior_state.get("workflow") == "reconciliation_review" and prior_state.get("batch_id"):
                    if self.service.batches.batch(prior_state["batch_id"])["status"] == "cancelled":
                        raise HTTPException(409, "The stopped financial batch is retained. Open its batch details before starting a new request.")
                    initial.update(workflow="reconciliation_review", batch_id=prior_state["batch_id"])
                initial.update(retry_of=retry_turn_id, required_tools=prior_state.get("required_tools", []),
                               steps=[{**step, "reused": True} for step in prior_state.get("steps", []) if step["status"] == "success"])
            with self.store.connect() as db:
                if not db.execute("SELECT id FROM data_turns WHERE chat_id=? LIMIT 1", (chat_id,)).fetchone():
                    db.execute("UPDATE data_chats SET title=? WHERE id=?", (instruction[:70], chat_id))
                cursor = db.execute("INSERT INTO data_turns(chat_id,instruction,status,state_json,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                                    (chat_id, instruction, "running", encode(initial), timestamp(), timestamp()))
                turn_id = cursor.lastrowid
            self.service.assistant_active = turn_id; self.service.cancel_signal.clear()
            self.service.future = self.service.executor.submit(self._work, chat_id, turn_id, instruction)
        return {"id": turn_id, "chat_id": chat_id, "status": "running"}

    def persist(self, chat_id, turn_id, state, turn, status):
        with self.store.connect() as db:
            db.execute("UPDATE data_chats SET state_json=?,updated_at=? WHERE id=?", (encode(state), timestamp(), chat_id))
            db.execute("UPDATE data_turns SET status=?,state_json=?,updated_at=? WHERE id=?", (status, encode(turn), timestamp(), turn_id))

    def _work(self, chat_id, turn_id, instruction):
        state = json.loads(self.chat(chat_id)["state_json"])
        turn = json.loads(self.store.one("SELECT state_json FROM data_turns WHERE id=?", (turn_id,))["state_json"])
        # Older retained frames predate difference metadata; recover it from actual successful calls.
        for step in turn["steps"]:
            info = state["frames"].get(step.get("result", {}).get("frame_id"))
            if not info: continue
            info["difference_columns"] = list(dict.fromkeys(info.get("difference_columns", []) + [column for column in info["columns"] if any(column in state["frames"].get(parent, {}).get("difference_columns", []) for parent in info.get("parents", []))]))
            args = step["arguments"]
            if step["tool"] == "derive_column" and args.get("operation") == "subtract" and args.get("column") in info["columns"]:
                info["difference_columns"] = list(dict.fromkeys(info["difference_columns"] + [args["column"]]))
            elif step["tool"] == "aggregate":
                inherited = state["frames"].get(args.get("frame_id"), {}).get("difference_columns", [])
                info["difference_columns"] = [metric["output"] for metric in args["metrics"] if metric["column"] in inherited and metric["op"] in ("sum", "mean", "min", "max")]
        tools = DataTools(self, chat_id, turn_id, state)
        status, failures, total_failures = "failed", 0, 0
        try:
            if review_requested(instruction):
                status = run_review(self, chat_id, turn_id, state, turn)
                return
            history = self.store.all("SELECT instruction,state_json FROM data_turns WHERE chat_id=? AND id<? ORDER BY id DESC LIMIT 4", (chat_id, turn_id))
            plan_schema = Plan.model_json_schema()
            plan_schema["properties"]["required_tools"]["items"] = {"type": "string", "enum": list(SPECS)}
            plan = Plan.model_validate(self.service.model.data_plan({"request": instruction,
                "latest_frame": state["latest_frame"], "frames": [{"id": key, "columns": info["columns"][:25]} for key, info in list(state["frames"].items())[-5:]],
                "files": tools.list_files()["files"][:30],
                "previous_request": history[0]["instruction"][:600] if history else None}, plan_schema))
            if any(tool not in SPECS for tool in plan.required_tools): raise ValueError("The model proposed an unsupported capability. Clarify the request.")
            if turn.get("retry_of"):
                plan.required_tools = list(dict.fromkeys(turn.get("required_tools", []) + plan.required_tools))
            turn["required_tools"] = plan.required_tools
            row_limit = explicit_row_limit(instruction)
            if row_limit is not None:
                if not 1 <= row_limit <= 5000: raise ValueError("Choose a row limit between 1 and 5,000.")
                if "limit_rows" not in plan.required_tools: plan.required_tools.append("limit_rows")
                turn["row_limit"] = row_limit
            if not plan.supported:
                turn["message"] = plan.message if not any("\u3400" <= char <= "\u9fff" for char in plan.message) else "Please describe a task using the available local data tools."
                status = "needs_input"
                return
            completed = {step["tool"] for step in turn["steps"] if step["status"] == "success"}
            pending = set(plan.required_tools) - completed
            action_budget = min(24, max(12, len(pending) * 2 + len(tools.list_files()["files"]) + 2))
            turn["action_budget"] = action_budget
            for action_index in range(action_budget):
                if self.service.cancel_signal.is_set(): status = "cancelled"; break
                context = {"request": instruction,
                           "previous_requests": [{"request": row["instruction"][:600], "response": json.loads(row["state_json"]).get("message", "")[:600]} for row in reversed(history)],
                           "files": tools.list_files()["files"][:30], "latest_frame": state["latest_frame"],
                           "frames": [{"id": info["id"], "rows": info["rows"], "columns": info["columns"][:25], "group": info.get("group", "input"), "operation": info["operation"]} for info in list(state["frames"].values())[-8:]],
                           "recent_steps": [{**step, "result": step.get("result") if len(encode(step.get("result"))) <= 2200 else {"summary": encode(step.get("result"))[:2000], "truncated": True}} for step in turn["steps"][-3:]],
                           "charts": list(state["charts"].keys())[-5:]}
                succeeded = {step["tool"] for step in turn["steps"] if step["status"] == "success"}
                context["remaining_required_tools"] = [tool for tool in plan.required_tools if tool not in succeeded]
                context["required_row_limit"] = row_limit
                context["actions_remaining"] = action_budget - action_index
                context["difference_chart_required"] = wants_difference_chart(instruction)
                context["_frame_catalog"] = [{"id": key, "columns": info["columns"], "rows": info["rows"], "difference_columns": info.get("difference_columns", [])} for key, info in state["frames"].items()]
                action = Action.model_validate(self.service.model.data_action(context, Action.model_json_schema()))
                if self.service.cancel_signal.is_set(): status = "cancelled"; break
                if action.action in ("final", "ask"):
                    if action.tool is not None or action.arguments: raise ValueError("A final response cannot include a tool call.")
                    missing = context["remaining_required_tools"]
                    if row_limit is not None:
                        latest = state["frames"].get(state["latest_frame"], {})
                        if latest.get("rows", row_limit + 1) > row_limit: missing = list(dict.fromkeys(missing + ["limit_rows"]))
                        # The newest export must reflect the final row-limited result.
                        exports = [step for step in turn["steps"] if step["tool"] == "export_report" and step["status"] == "success"]
                        if exports and (state["latest_frame"] not in exports[-1]["result"].get("frame_ids", []) or any(self.state_rows(state, frame_id) > row_limit for frame_id in exports[-1]["result"].get("frame_ids", []))):
                            missing = list(dict.fromkeys(missing + ["export_report"]))
                    if "create_chart" in plan.required_tools and "export_report" in plan.required_tools:
                        exported_charts = {chart for step in turn["steps"] if step["tool"] == "export_report" and step["status"] == "success" for chart in step["result"].get("chart_ids", [])}
                        created_charts = {step["result"]["chart_id"] for step in turn["steps"] if step["tool"] == "create_chart" and step["status"] == "success"}
                        if created_charts and not created_charts.issubset(exported_charts): missing = list(dict.fromkeys(missing + ["export_report"]))
                    if action.action == "final" and (missing or turn["steps"] and turn["steps"][-1]["status"] == "error"):
                        failures += 1
                        total_failures += 1
                        turn["steps"].append({"tool": "completion_check", "arguments": {}, "status": "error", "error": "The task is not complete. Still required: " + ", ".join(missing) + ". Correct the last error and call the remaining tools. A written claim is not execution."})
                        self.persist(chat_id, turn_id, state, turn, "running")
                        if failures >= 3 or total_failures >= 8: raise ValueError("The requested outputs were not produced. Retained results and tool errors are available; continue with a smaller request.")
                        continue
                    if action.action == "final" and not turn["steps"] and not state["frames"]:
                        raise ValueError("No data has been inspected. Ask a question or inspect a table first.")
                    message = action.message.strip()
                    if not message or any("\u3400" <= char <= "\u9fff" for char in message):
                        message = "Please clarify the file, columns or operation to use." if action.action == "ask" else "The requested data steps have finished. See the results and tool history below."
                    turn["message"] = message
                    if action.action == "final" and context["difference_chart_required"]:
                        charts = [step["result"] for step in turn["steps"] if step["tool"] == "create_chart" and step["status"] == "success"]
                        if charts:
                            chart = charts[-1]; summary = tools.frame(chart["frame_id"])
                            lines = ["Completed the requested data steps.", f"Chart: {chart['y']} by {chart['x']}.", "Computed summary:"]
                            for row in summary.head(10).to_dict(orient="records"):
                                lines.append("; ".join(f"{key}: {value}" for key, value in list(row.items())[:8]))
                            if len(summary) > 10: lines.append(f"Showing 10 of {len(summary)} groups.")
                            if any(step["tool"] == "export_report" and step["status"] == "success" for step in turn["steps"]): lines.append("The generated Excel report is available below.")
                            turn["model_message"] = message
                            turn["message"] = "\n".join(lines)
                    # Keep currencies grounded in explicit user/source metadata.
                    known = {currency for info in state["frames"].values() for currency in info.get("currencies", [])}
                    for symbol, code in (("$", "USD"), ("£", "GBP"), ("€", "EUR")):
                        if symbol not in instruction and code not in known and code not in instruction.upper():
                            turn["message"] = re.sub(re.escape(symbol) + r"(?=\s*\d)", "", turn["message"])
                    question = message.endswith("?") or re.search(r"\b(please confirm|need your confirmation|shall I proceed|may I inspect)\b", message, re.I)
                    status = "needs_input" if action.action == "ask" or question else "completed"
                    break
                try:
                    if action.tool == "limit_rows" and row_limit is not None and "sort_rows" in plan.required_tools:
                        sorted_frames = {step["result"]["frame_id"] for step in turn["steps"] if step["tool"] == "sort_rows" and step["status"] == "success"}
                        if action.arguments.get("frame_id") not in sorted_frames:
                            raise ValueError("Sort the complete candidate table before taking its top rows. Call sort_rows first, then pass its returned frame_id to limit_rows.")
                    if action.tool == "create_chart" and context["difference_chart_required"]:
                        info = state["frames"].get(action.arguments.get("frame_id"), {})
                        if action.arguments.get("y") not in info.get("difference_columns", []):
                            raise ValueError("The requested chart must plot a computed amount difference, not order counts. Aggregate the detail difference with op=sum into the summary or subtract the two summary amount totals, then plot that difference column.")
                    result = tools.call(action.tool, action.arguments)
                    failures = 0
                    # Persist actual tool results, never a model's claimed execution.
                    result_text = encode(result)
                    if len(result_text) > 6500:
                        result = {"summary": result_text[:6000], "truncated": True, "note": "Use selected columns or a smaller table for detail."}
                    turn["steps"].append({"tool": action.tool, "arguments": action.arguments, "status": "success", "result": result})
                    self.persist(chat_id, turn_id, state, turn, "running")
                except (ValueError, TypeError, KeyError, ArithmeticError) as exc:
                    failures += 1
                    total_failures += 1
                    turn["steps"].append({"tool": action.tool, "arguments": action.arguments, "status": "error", "error": str(exc)[:1200]})
                    self.persist(chat_id, turn_id, state, turn, "running")
                    if failures >= 3: raise ValueError("Three tool calls failed consecutively. Saved results are available; retry the request or clarify the last error.") from exc
                    if total_failures >= 8: raise ValueError("The tool recovery limit was reached. Saved results are available; retry with a smaller request.") from exc
            else:
                turn["error"] = f"The {action_budget}-step limit was reached. Saved results are available; retry or continue with a smaller request."
            if status == "cancelled": turn["message"] = "Stopped at a tool boundary. Existing results and reports were retained."
        except Exception as exc:
            turn["error"] = str(exc)[:1500]
        finally:
            try: self.persist(chat_id, turn_id, state, turn, status)
            finally:
                with self.service.guard: self.service.assistant_active = None

    def cancel(self, turn_id):
        with self.service.guard:
            if self.service.assistant_active != turn_id: raise HTTPException(409, "This request is not running.")
            self.service.cancel_signal.set()
        return {"cancel_requested": True}

    @staticmethod
    def state_rows(state, frame_id):
        return state["frames"].get(frame_id, {}).get("rows", 50001)

    def artifact(self, artifact_id):
        row = self.store.one("SELECT * FROM data_artifacts WHERE id=?", (artifact_id,))
        if not row: raise HTTPException(404, "Report was not found.")
        path = self.service.path(row["path"])
        if not path.is_file() or digest(path) != row["sha256"]: raise HTTPException(409, "The report failed its integrity check.")
        return path
