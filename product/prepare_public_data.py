"""Prepare attributed UCI benchmark files. Never import or seed application data."""
from collections import OrderedDict
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import argparse
import csv
import gzip
import hashlib
import io
import json
from pathlib import Path
import urllib.request
import zipfile

from openpyxl import load_workbook

URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"
WORKBOOK_URL = "https://archive.ics.uci.edu/ml/machine-learning-databases/00352/Online%20Retail.xlsx"
CSV_URL = "https://raw.githubusercontent.com/elastic/eland/main/docs/sphinx/examples/data/online-retail.csv.gz"
SOURCE = "https://archive.ics.uci.edu/dataset/352/online+retail"
HEADERS = ["supplier_id", "supplier_name", "order_id", "amount", "date", "payment_status", "record_status"]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--archive", type=Path, help="Use an already downloaded original ZIP")
    parser.add_argument("--workbook", type=Path, help="Use the publisher's original Excel workbook instead of the ZIP")
    parser.add_argument("--csv-gzip", type=Path, help="Use the attributed Elastic Eland public CSV sample mirror")
    parser.add_argument("--output", type=Path, default=Path(__file__).parent / "examples")
    options = parser.parse_args()
    use_csv = bool(options.csv_gzip or not (options.archive or options.workbook))
    options.output.mkdir(parents=True, exist_ok=True)
    if sum(bool(value) for value in (options.archive, options.workbook, options.csv_gzip)) > 1:
        parser.error("Choose one input format.")
    if options.csv_gzip:
        content = options.csv_gzip.read_bytes()
    elif options.workbook:
        content = options.workbook.read_bytes()
    elif options.archive:
        content = options.archive.read_bytes()
    else:
        with urllib.request.urlopen(CSV_URL, timeout=120) as response:
            content = response.read(2 * 1024 * 1024)
    workbook = None
    if use_csv:
        rows = csv.reader(io.StringIO(gzip.decompress(content).decode("utf-8-sig")))
    elif options.workbook:
        workbook_bytes = content
    else:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            name = next(name for name in archive.namelist() if name.endswith(".xlsx"))
            workbook_bytes = archive.read(name)
    if not use_csv:
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
    try:
        headers = next(rows)
        groups = OrderedDict()
        for values in rows:
            item = dict(zip(headers, values))
            invoice, customer = str(item["InvoiceNo"]), item["CustomerID"]
            if customer in (None, "", "nan") or invoice.upper().startswith("C") or Decimal(str(item["Quantity"])) <= 0 or Decimal(str(item["UnitPrice"])) <= 0:
                continue
            customer = str(int(Decimal(str(customer))))
            key = (invoice, customer)
            if key not in groups:
                if len(groups) >= 100:
                    continue
                invoice_date = item["InvoiceDate"]
                if isinstance(invoice_date, str):
                    invoice_date = datetime.fromisoformat(invoice_date)
                groups[key] = {"supplier_id": "PUBLIC-CUSTOMER-" + customer,
                               "supplier_name": "Public benchmark counterparty " + customer,
                               "order_id": "UCI-" + invoice, "amount": Decimal("0"),
                               "date": invoice_date.date().isoformat(),
                               "payment_status": "pending", "record_status": "active"}
            groups[key]["amount"] += Decimal(str(item["UnitPrice"])) * Decimal(str(item["Quantity"]))
    finally:
        if workbook:
            workbook.close()
    reference = list(groups.values())
    for record in reference:
        record["amount"] = str(record["amount"].quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    source = [dict(row) for row in reference]
    if len(source) < 3:
        raise ValueError("The public dataset did not contain enough complete positive invoices.")
    source[0]["amount"] = str(Decimal(source[0]["amount"]) + Decimal("25.00"))
    source[1]["payment_status"] = "paid"
    source[2]["supplier_name"] = ""
    for filename, records in (("public_reference.csv", reference), ("public_orders_with_findings.csv", source)):
        with (options.output / filename).open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerows(records)
    provenance = {"title": "UCI Online Retail-derived reconciliation benchmark",
                  "citation": "Chen, D. (2015). Online Retail [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5BW33",
                  "source": SOURCE, "download": CSV_URL if use_csv else WORKBOOK_URL if options.workbook else URL,
                  "download_scope": "Elastic Eland's 15,000-line public sample" if use_csv else "Original UCI workbook",
                  "license": "CC BY 4.0", "license_url": "https://creativecommons.org/licenses/by/4.0/",
                  "original_download_sha256": hashlib.sha256(content).hexdigest(),
                  "rows": len(reference), "currency": "GBP (sterling), unchanged; no conversion",
                  "transformations": ["Aggregate positive, non-cancelled invoice line totals by InvoiceNo and CustomerID, using Decimal.",
                                      "Use CustomerID as an illustrative counterparty surrogate in supplier_id; this is retail customer data, not real supplier data.",
                                      "Create surrogate English counterparty names; convert invoice date to YYYY-MM-DD.",
                                      "Construct pending payment_status and active record_status; these statuses are not present in UCI.",
                                      "Derive the reference from the same public dataset; it is not independent enterprise BI evidence."],
                  "injected_findings": [{"working_row": 2, "change": "amount +25 GBP"},
                                        {"working_row": 3, "change": "payment_status pending -> paid"},
                                        {"working_row": 4, "change": "supplier_name removed"}]}
    (options.output / "PROVENANCE.json").write_text(json.dumps(provenance, indent=2), encoding="utf-8")
    print(f"Prepared {len(reference)} public benchmark orders and three labeled findings in {options.output}")


if __name__ == "__main__":
    main()
