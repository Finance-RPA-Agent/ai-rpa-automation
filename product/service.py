"""Persistent product workflow over the committed AgentRuntime and four tools.

Each job has its own financial database and immutable reference snapshot. This
prevents one BI import from changing another job's evidence or archive receipts.
"""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from decimal import Decimal, InvalidOperation
from functools import wraps
from pathlib import Path
import copy
import csv
import hashlib
import json
import secrets
import threading
import zipfile

from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend.app.config import Settings
from backend.app.database import Base
from backend.app.models import FinancialRecord, OaReconciliationRecord, Task
from backend.app.services.agent_runtime import AgentRuntime
from backend.app.tools import read_data_file, validate_data
from backend.app.tools.rules import REQUIRED_FIELDS, format_amount, normalise_identifier, parse_amount
from product.model import OllamaProvider
from product.storage import Store, digest, encode, timestamp

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
FIELDS = list(REQUIRED_FIELDS)


def serialized(method):
    @wraps(method)
    def call(self, *args, **kwargs):
        with self.guard:
            return method(self, *args, **kwargs)
    return call


class Cancelled(Exception):
    pass


def result(envelope):
    if envelope.get("status") != "success":
        raise ValueError(f"{envelope.get('error_code', 'READ_ERROR')}: {envelope.get('message', 'Operation failed.')}")
    return envelope["result"]


def write_csv(path, columns, rows, *, spreadsheet_safe=False):
    def safe(value):
        # Neutralize formula-bearing text only in user-facing exports.
        if spreadsheet_safe and isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + value
        return value
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: safe(value) for key, value in row["data"].items()} for row in rows)


class ProductService:
    def __init__(self, directory, *, model=None):
        self.store = Store(directory)
        self.model = model or OllamaProvider()
        self.settings = Settings(self.store.root, self.store.root / "files", self.store.root / "jobs", self.store.root / "exports")
        self.executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="finance-flow")
        self.guard = threading.RLock()
        self.active = None
        self.batch_active = None
        self.assistant_active = None
        self.cancel_signal = threading.Event()
        self.future = None
        with self.store.connect() as db:
            rows = db.execute("SELECT * FROM jobs WHERE status='running' OR json_extract(state_json,'$.plan_status')='generating'").fetchall()
            for row in rows:
                state = json.loads(row["state_json"])
                state.update(error="The application stopped during execution. Existing records were retained; retry to continue.", plan_status="interrupted")
                db.execute("UPDATE jobs SET status='failed',state_json=?,updated_at=? WHERE id=?", (encode(state), timestamp(), row["id"]))
            db.execute("UPDATE batches SET status='failed',error=?,updated_at=? WHERE status IN ('scanning','running')",
                       ("The application stopped. Captured files and existing receipts were retained. Review before retrying.", timestamp()))
        from product.batch import FolderBatches
        self.batches = FolderBatches(self)
        from product.commands import Commands
        self.commands = Commands(self)
        from product.assistant import DataAssistant
        self.assistant = DataAssistant(self)

    def path(self, relative):
        path = (self.store.root / relative).resolve()
        if not path.is_relative_to(self.store.root):
            raise ValueError("Stored path is outside application data.")
        return path

    def file(self, file_id):
        row = self.store.one("SELECT * FROM files WHERE id=?", (file_id,))
        if not row:
            raise HTTPException(404, "File was not found.")
        return row

    def load_file(self, file_id):
        file = self.file(file_id)
        if digest(self.path(file["raw_path"])) != file["sha256"] or digest(self.path(file["canonical_path"])) != file["canonical_sha256"]:
            raise ValueError("A stored file changed outside the application. Import a new immutable version.")
        return result(read_data_file(str(self.path(file["canonical_path"])), settings=self.settings))

    def upload(self, name, content, role):
        if role not in ("source", "reference"):
            raise HTTPException(422, "File role must be source or reference.")
        suffix = Path(name or "").suffix.lower()
        if suffix not in (".xlsx", ".csv"):
            raise HTTPException(422, "Only .xlsx and UTF-8 .csv files are supported.")
        if not content or len(content) > MAX_BYTES:
            raise HTTPException(422, "Choose a nonempty file of at most 10 MB.")
        file_id = secrets.token_hex(16)
        folder = self.store.root / "files" / file_id
        folder.mkdir()
        raw = folder / ("original" + suffix)
        raw.write_bytes(content)
        if suffix == ".csv":
            try:
                with raw.open(encoding="utf-8-sig", newline="") as handle:
                    count = 0
                    for values in csv.reader(handle, strict=True):
                        if not values:
                            continue
                        if len(values) > 200:
                            raise ValueError("Files may contain at most 200 columns.")
                        count += 1
                        if count > MAX_ROWS + 1:
                            raise ValueError(f"Files may contain at most {MAX_ROWS:,} data rows.")
            except (UnicodeDecodeError, csv.Error) as exc:
                raise ValueError("The CSV must use UTF-8 encoding and correctly closed quoted fields.") from exc
        if suffix == ".xlsx":
            try:
                with zipfile.ZipFile(raw) as archive:
                    if sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                        raise ValueError("The workbook's decompressed content exceeds 100 MB.")
            except zipfile.BadZipFile as exc:
                raise ValueError("This is not a valid .xlsx workbook.") from exc
            from openpyxl import load_workbook
            try:
                workbook = load_workbook(raw, read_only=True, data_only=True)
            except Exception as exc:
                raise ValueError("The Excel workbook could not be opened. Check that it is a valid, unencrypted .xlsx file.") from exc
            try:
                if not workbook.worksheets:
                    raise ValueError("The Excel workbook has no readable worksheets.")
                sheet = workbook.worksheets[0]
                if (sheet.max_row and sheet.max_row > MAX_ROWS + 1) or (sheet.max_column and sheet.max_column > 200):
                    raise ValueError(f"The first worksheet must fit within {MAX_ROWS:,} data rows and 200 columns.")
            finally:
                workbook.close()
        table = result(read_data_file(str(raw), settings=self.settings))
        if not 0 < table["row_count"] <= MAX_ROWS:
            raise ValueError(f"Files must contain 1–{MAX_ROWS:,} data rows.")
        if table["column_issues"]:
            raise ValueError("Headers are empty or duplicated. Correct the column names before importing.")
        mapping = {field: field for field in FIELDS if field in table["columns"]}
        canonical = raw
        with self.store.connect() as db:
            db.execute("INSERT INTO files VALUES(?,?,?,?,?,?,?,?,?,?,?)", (
                file_id, Path(name).name, role, str(raw.relative_to(self.store.root)), str(canonical.relative_to(self.store.root)),
                digest(raw), digest(canonical), encode(table["columns"]), encode(mapping), table["row_count"], timestamp(),
            ))
        return self.file_view(file_id)

    def file_view(self, file_id, *, preview=True):
        row = self.file(file_id)
        table = self.load_file(file_id) if preview else None
        return {"id": row["id"], "name": row["name"], "role": row["role"], "row_count": row["row_count"],
                "columns": json.loads(row["columns_json"]), "mapping": json.loads(row["mapping_json"]),
                "ready": set(json.loads(row["mapping_json"])) == set(FIELDS),
                "sha256": row["sha256"], "created_at": row["created_at"], "preview": table["records"][:5] if table else []}

    @serialized
    def map_file(self, file_id, mapping):
        row = self.file(file_id)
        if self.store.one("SELECT id FROM jobs WHERE source_id=? OR reference_id=? LIMIT 1", (file_id, file_id)):
            raise HTTPException(409, "This file is already bound to a task. Import a new version to change its mapping.")
        columns = json.loads(row["columns_json"])
        if set(mapping) != set(FIELDS) or len(set(mapping.values())) != len(FIELDS) or any(value not in columns for value in mapping.values()):
            raise HTTPException(422, "Map each required field to a different existing column.")
        if digest(self.path(row["raw_path"])) != row["sha256"]:
            raise HTTPException(409, "The original file changed; please import it again.")
        table = result(read_data_file(str(self.path(row["raw_path"])), settings=self.settings))
        records = [{"row": index + 2, "data": {field: item["data"].get(column) for field, column in mapping.items()}}
                   for index, item in enumerate(table["records"])]
        target = self.path(row["raw_path"]).parent / (secrets.token_hex(8) + ".csv")
        write_csv(target, FIELDS, records)
        if row["role"] == "reference":
            checked = result(validate_data(str(target), settings=self.settings))
            if checked["issues"]:
                raise ValueError("The BI reference contains invalid or duplicate data: " + checked["issues"][0]["message"])
        with self.store.connect() as db:
            db.execute("UPDATE files SET canonical_path=?,canonical_sha256=?,mapping_json=? WHERE id=?",
                       (str(target.relative_to(self.store.root)), digest(target), encode(mapping), file_id))
        return self.file_view(file_id)

    def get_job(self, job_id):
        job = self.store.one("SELECT * FROM jobs WHERE id=?", (job_id,))
        if not job:
            raise HTTPException(404, "Task was not found.")
        return job, json.loads(job["state_json"])

    @serialized
    def create(self, title, instruction, source_id, reference_id, ai):
        source, reference = self.file(source_id), self.file(reference_id)
        if source["role"] != "source" or reference["role"] != "reference" or source_id == reference_id:
            raise HTTPException(422, "Choose a source file and a separate BI reference import.")
        for file in (source, reference):
            if set(json.loads(file["mapping_json"])) != set(FIELDS):
                raise HTTPException(422, "Complete both files' column mappings before creating the task.")
            self.load_file(file["id"])
        checked = result(validate_data(str(self.path(reference["canonical_path"])), settings=self.settings))
        if checked["issues"]:
            raise ValueError("The BI reference must pass field validation: " + checked["issues"][0]["message"])
        state = {"working_file": source["canonical_path"], "working_hash": source["canonical_sha256"],
                 "plan": None, "plan_status": "not_generated", "ai": ai,
                 "records": [], "issues": [], "archive_count": 0, "error": None, "stage": "Ready to generate plan",
                 "initial_issues": 0, "evidence_digest": None, "trace": [], "cancel_requested": False}
        with self.store.connect() as db:
            cursor = db.execute("INSERT INTO jobs(title,instruction,source_id,reference_id,status,state_json,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                                (title, instruction, source_id, reference_id, "pending", encode(state), timestamp(), timestamp()))
            job_id = cursor.lastrowid
        self.store.event(job_id, "Task created", "Source and BI reference snapshots are attached.")
        return self.detail(job_id)

    def schedule(self, job_id, operation, actor):
        with self.guard:
            if self.active is not None or self.batch_active is not None or self.assistant_active is not None:
                raise HTTPException(409, "Another task is using the local worker. Wait for it to finish.")
            job, state = self.get_job(job_id)
            if job["status"] in ("completed", "cancelled", "running"):
                raise HTTPException(409, "This task is running or has already ended.")
            if operation == "run":
                if not state["plan"] or state["plan_status"] != "ready":
                    raise HTTPException(409, "Generate and review the execution plan first.")
                if state["issues"]:
                    raise HTTPException(409, "Resolve the remaining findings before continuing.")
                state.update(error=None, cancel_requested=False, stage="Starting reconciliation")
                self.store.save_job(job, state, "running")
            else:
                if job["status"] == "waiting_confirmation":
                    raise HTTPException(409, "Review the existing task before changing its plan.")
                state.update(plan_status="generating", error=None, stage="Generating local plan")
                self.store.save_job(job, state, "pending")
            self.active = job_id
            self.cancel_signal.clear()
            self.future = self.executor.submit(self._work, job_id, operation, actor)
        return self.detail(job_id)

    def _work(self, job_id, operation, actor):
        try:
            if operation == "plan":
                job, state = self.get_job(job_id)
                if state["ai"]:
                    plan = self.model.plan(job["instruction"], str(self.path(state["working_file"])))
                else:
                    plan = {"mode": "deterministic", "model": None,
                            "summary": "Use the reviewed supplier reconciliation workflow. Model assistance is explicitly disabled for this task.",
                            "steps": ["Read source file", "Validate fields", "Query imported BI reference", "Compare sources", "Human review if needed", "Archive clean results"]}
                if self.cancel_signal.is_set():
                    raise Cancelled()
                job, state = self.get_job(job_id)
                state.update(plan=plan, plan_status="ready", stage="Plan ready for approval")
                self.store.save_job(job, state, "pending")
                self.store.event(job_id, "Plan generated", plan["summary"], details={"mode": plan["mode"], "model": plan["model"]})
            else:
                self.execute(job_id, actor)
        except Cancelled:
            job, state = self.get_job(job_id)
            state.update(stage="Cancelled", plan_status="ready" if state["plan"] else "not_generated", cancel_requested=False)
            self.store.save_job(job, state, "cancelled")
            self.store.event(job_id, "Task cancelled", "Completed operations and audit records were retained.", "cancelled")
        except Exception as exc:
            job, state = self.get_job(job_id)
            state.update(error=str(exc)[:1500], stage="Execution stopped", plan_status="ready" if state["plan"] else "failed")
            self.store.save_job(job, state, "failed")
            self.store.event(job_id, "Operation failed", str(exc)[:1500], "failed")
        finally:
            with self.guard:
                self.active = None

    def job_database(self, job):
        folder = self.store.root / "jobs" / str(job["id"])
        folder.mkdir(exist_ok=True)
        engine = create_engine(f"sqlite:///{(folder / 'financial.sqlite').as_posix()}", connect_args={"check_same_thread": False})
        Base.metadata.create_all(engine)
        session = sessionmaker(engine, expire_on_commit=False)()
        if not session.get(Task, job["id"]):
            session.add(Task(id=job["id"], description=job["instruction"], status="pending"))
            reference = self.load_file(job["reference_id"])
            for item in reference["records"]:
                values = dict(item["data"])
                values["supplier_id"] = normalise_identifier(values["supplier_id"])
                values["order_id"] = normalise_identifier(values["order_id"])
                values["amount"] = format_amount(parse_amount(values["amount"]))
                session.add(FinancialRecord(**{field: str(values[field]).strip() for field in FIELDS}, source="user_imported_bi_export"))
            session.commit()
        return session, engine

    def execute(self, job_id, actor):
        job, state = self.get_job(job_id)
        self.load_file(job["source_id"])
        self.load_file(job["reference_id"])
        working = self.path(state["working_file"])
        if digest(working) != state["working_hash"]:
            raise ValueError("The task's working revision changed outside the application.")
        db, engine = self.job_database(job)
        service = self
        self.store.event(job_id, "Execution started", "The approved plan is running against immutable local file revisions.", details={"actor": actor, "revision": job["revision"]})

        class TrackedRuntime(AgentRuntime):
            current_trace = None

            def _execute(runtime, tool, arguments, trace):
                runtime.current_trace = trace
                if service.cancel_signal.is_set():
                    raise Cancelled()
                result_value, decision = super()._execute(tool, arguments, trace)
                current, saved = service.get_job(job_id)
                labels = {"read_data_file": "Read source file", "validate_data": "Validate and compare", "query_bi": "Query BI reference", "submit_to_oa": "Archive reconciled record"}
                saved["stage"] = labels.get(str(tool), str(tool))
                if str(tool) == "read_data_file":
                    saved["records"] = result_value["records"]
                if str(tool) == "validate_data":
                    saved["issues"] = result_value["issues"]
                service.store.save_job(current, saved, "running")
                service.store.event(job_id, saved["stage"], decision.reason,
                                    details={"tool": str(tool), "classification": decision.classification.value,
                                             "attempts": trace[-1].attempt + 1})
                return result_value, decision

        runtime = TrackedRuntime(db=db, settings=self.settings)
        try:
            outcome = runtime.run({"task_id": job_id, "file_path": str(working),
                                   "operation_id": f"task-{job_id}-revision-{job['revision']}"})
            self.persist_trace(job_id, outcome.trace, db)
            job, state = self.get_job(job_id)
            state["evidence_digest"] = outcome.confirmation_request.evidence_digest if outcome.confirmation_request else None
            state["issues"] = self.enrich_issues(state["issues"], state["records"], db)
            state["initial_issues"] = max(state["initial_issues"], len(state["issues"]))
            state["error"] = outcome.reason if outcome.error_code else None
            status = outcome.task_status.value
            state["stage"] = {"completed": "Reconciliation complete", "waiting_confirmation": "Human review required", "failed": "Execution failed"}.get(status, status)
            if status == "completed":
                output = self.store.root / "exports" / f"FF-{job_id:05d}-reviewed.csv"
                write_csv(output, FIELDS, state["records"], spreadsheet_safe=True)
                state["export"] = str(output.relative_to(self.store.root))
            self.store.save_job(job, state, status)
            self.store.event(job_id, state["stage"], outcome.reason, status)
        except BaseException:
            self.persist_trace(job_id, runtime.current_trace or [], db)
            raise
        finally:
            db.close()
            engine.dispose()

    def persist_trace(self, job_id, trace, db):
        job, state = self.get_job(job_id)
        state["trace"] = [asdict(attempt) for attempt in trace]
        for receipt in db.query(OaReconciliationRecord).order_by(OaReconciliationRecord.id):
            record_id = f"FF-{job_id:05d}-{receipt.id:05d}"
            with self.store.connect() as central:
                central.execute("INSERT OR IGNORE INTO archive VALUES(?,?,?,?,?,?)", (
                    record_id, job_id, receipt.operation_id, receipt.record_id, receipt.payload, timestamp()))
        state["archive_count"] = self.store.one("SELECT COUNT(*) AS n FROM archive WHERE job_id=?", (job_id,))["n"]
        self.store.save_job(job, state)

    @staticmethod
    def enrich_issues(issues, records, db):
        by_row = {row["row"]: row["data"] for row in records}
        enriched = []
        for issue in issues:
            original = by_row.get(issue.get("row"), {})
            matches = db.query(FinancialRecord).filter_by(
                supplier_id=normalise_identifier(original.get("supplier_id")),
                order_id=normalise_identifier(original.get("order_id")),
            ).all() if original else []
            field = issue.get("field")
            reference = getattr(matches[0], field, None) if len(matches) == 1 and field in FIELDS else None
            allowed = field in {"supplier_name", "amount", "date", "payment_status", "record_status"}
            identity = encode([issue.get("row"), field, issue["type"]])
            enriched.append(dict(issue, id=hashlib.sha256(identity.encode()).hexdigest()[:16],
                                 original_value=original.get(field), reference_value=reference,
                                 supplier_id=original.get("supplier_id"), order_id=original.get("order_id"),
                                 can_correct=bool(allowed and reference is not None)))
        return enriched

    def review(self, job_id, revision, evidence, changes, reason, actor):
        with self.guard:
            if self.active is not None or self.batch_active is not None or self.assistant_active is not None:
                raise HTTPException(409, "Wait for the active operation to finish.")
            job, state = self.get_job(job_id)
            if job["status"] != "waiting_confirmation" or revision != job["revision"] or evidence != state["evidence_digest"]:
                raise HTTPException(409, "This review is stale. Refresh the current task evidence.")
            self.load_file(job["source_id"])
            self.load_file(job["reference_id"])
            working = self.path(state["working_file"])
            if digest(working) != state["working_hash"]:
                raise HTTPException(409, "The working file changed. Review cannot proceed.")
            if not changes:
                raise HTTPException(422, "Select at least one correction.")
            current = {issue["id"]: issue for issue in state["issues"]}
            if len(changes) != len(set(item["issue_id"] for item in changes)):
                raise HTTPException(422, "A finding cannot be corrected twice in one decision.")
            records = copy.deepcopy(state["records"])
            by_row = {row["row"]: row["data"] for row in records}
            audit = []
            for change in changes:
                issue = current.get(change["issue_id"])
                if not issue or not issue["can_correct"]:
                    raise HTTPException(422, "A selected finding requires a corrected source file, not automatic resolution.")
                value = issue["reference_value"] if change["mode"] == "reference" else change["value"]
                if value is None or not str(value).strip():
                    raise HTTPException(422, "A correction must include a nonempty replacement value.")
                by_row[issue["row"]][issue["field"]] = value
                audit.append({"source_row": issue["row"], "field": issue["field"], "before": issue["original_value"],
                              "after": value, "mode": change["mode"], "issue_type": issue["type"]})
            new_revision = job["revision"] + 1
            target = self.store.root / "jobs" / str(job_id) / f"revision-{new_revision}-{secrets.token_hex(4)}.csv"
            write_csv(target, FIELDS, records)
            table = result(read_data_file(str(target), settings=self.settings))
            reference = self.load_file(job["reference_id"])
            checked = result(validate_data(str(target), reference_records=[r["data"] for r in reference["records"]], settings=self.settings))
            db, engine = self.job_database(job)
            try:
                new_issues = self.enrich_issues(checked["issues"], table["records"], db)
            finally:
                db.close()
                engine.dispose()
            state.update(working_file=str(target.relative_to(self.store.root)), working_hash=digest(target),
                         records=table["records"], issues=new_issues, evidence_digest=digest(target),
                         stage="Corrections saved; ready to resume" if not new_issues else "Further review required")
            with self.store.connect() as db:
                db.execute("INSERT INTO decisions(job_id,revision,evidence_digest,actor,changes_json,reason,created_at) VALUES(?,?,?,?,?,?,?)",
                           (job_id, revision, evidence, actor, encode(audit), reason.strip(), timestamp()))
                db.execute("UPDATE jobs SET revision=?,state_json=?,updated_at=? WHERE id=?", (new_revision, encode(state), timestamp(), job_id))
                db.execute("INSERT INTO events(job_id,action,status,message,details_json,created_at) VALUES(?,?,?,?,?,?)",
                           (job_id, "Human correction recorded", "reviewed", reason.strip() or "The reviewer confirmed the selected corrections.", encode(audit), timestamp()))
            return self.detail(job_id)

    def cancel(self, job_id, actor):
        with self.guard:
            if self.batch_active is not None or self.assistant_active is not None:
                raise HTTPException(409, "Use the folder batch cancellation control while a batch is active.")
            job, state = self.get_job(job_id)
            if job["status"] in ("completed", "cancelled"):
                raise HTTPException(409, "This task has already ended.")
            if self.active == job_id:
                self.cancel_signal.set()
                state["cancel_requested"] = True
                self.store.save_job(job, state)
                self.store.event(job_id, "Cancellation requested", "Execution will stop at the next safe tool boundary.", details={"actor": actor})
            else:
                self.store.save_job(job, state, "cancelled")
                self.store.event(job_id, "Task cancelled", "Existing audit and archive records were retained.", "cancelled", {"actor": actor})
        return self.detail(job_id)

    def detail(self, job_id, *, full=False, batch_id=None):
        job, state = self.get_job(job_id)
        source, reference = self.file(job["source_id"]), self.file(job["reference_id"])
        total = Decimal("0")
        for row in state["records"]:
            try:
                total += parse_amount(row["data"].get("amount"))
            except ValueError:
                pass
        summary = {key: job[key] for key in ("id", "title", "instruction", "status", "revision", "created_at", "updated_at")}
        summary.update(code=f"FF-{job_id:05d}", source_name=source["name"], reference_name=reference["name"],
                       row_count=len(state["records"]), issue_count=len(state["issues"]), amount=format_amount(total),
                       archive_count=state["archive_count"], stage=state["stage"], error=state["error"],
                       plan=state["plan"], plan_status=state["plan_status"], ai=state["ai"],
                       cancel_requested=state["cancel_requested"], evidence_digest=state["evidence_digest"],
                       source_sha256=source["sha256"], reference_sha256=reference["sha256"])
        if full:
            summary.update(records=state["records"], issues=state["issues"], trace=state["trace"])
        else:
            summary.update(records=state["records"][:100], issues=state["issues"][:200],
                           truncated_preview=len(state["records"]) > 100, truncated_issues=len(state["issues"]) > 200)
        summary["events"] = self.store.all("SELECT * FROM events WHERE job_id=? ORDER BY id DESC LIMIT 200", (job_id,))
        decisions = self.store.all("SELECT * FROM decisions WHERE job_id=? ORDER BY id", (job_id,))
        summary["decisions"] = [dict(row, changes=json.loads(row["changes_json"])) for row in decisions]
        summary["archive"] = [dict(row, data=json.loads(row["payload_json"])) for row in self.store.all("SELECT * FROM archive WHERE job_id=? ORDER BY id", (job_id,))]
        origin = self.store.one("SELECT batch_id,rows_json FROM batch_jobs WHERE job_id=?" +
                                (" AND batch_id=?" if batch_id is not None else "") + " ORDER BY batch_id LIMIT 1",
                                (job_id, batch_id) if batch_id is not None else (job_id,))
        if batch_id is not None and not origin:
            raise HTTPException(404, "This task is not associated with the selected folder batch.")
        if origin:
            origins = json.loads(origin["rows_json"])
            summary["batch_id"] = origin["batch_id"]
            for row in summary.get("records", []):
                if 0 <= row["row"] - 2 < len(origins):
                    row["original_row"] = origins[row["row"] - 2]
            for issue in summary.get("issues", []):
                index = (issue.get("row") or 0) - 2
                if 0 <= index < len(origins):
                    issue["original_row"] = origins[index]
            for decision in summary["decisions"]:
                for change in decision["changes"]:
                    index = change["source_row"] - 2
                    if 0 <= index < len(origins):
                        change["original_row"] = origins[index]
        return summary

    def tasks(self):
        items = []
        for row in self.store.all("SELECT id FROM jobs ORDER BY id DESC"):
            task = self.detail(row["id"])
            for key in ("records", "issues", "events", "decisions", "archive"):
                task.pop(key, None)
            items.append(task)
        return items

    def backup(self):
        with self.guard:
            if self.active is not None or self.batch_active is not None or self.assistant_active is not None:
                raise HTTPException(409, "Wait for the active operation before backing up.")
            name = "finance-flow-backup-" + timestamp().replace(":", "-") + ".zip"
            target = self.store.root / "backups" / name
            snapshot = self.store.root / "backups" / (secrets.token_hex(8) + ".sqlite")
            self.store.backup(snapshot)
            manifest = {}
            with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.write(snapshot, "finance-flow.sqlite")
                manifest["finance-flow.sqlite"] = digest(snapshot)
                for folder in ("files", "jobs", "exports"):
                    for path in (self.store.root / folder).rglob("*"):
                        if path.is_file() and not path.name.endswith(("-wal", "-shm", "-journal")):
                            relative = path.relative_to(self.store.root).as_posix()
                            archive.write(path, relative)
                            manifest[relative] = digest(path)
                archive.writestr("manifest.json", encode({"version": 1, "files": manifest}))
            snapshot.unlink()
            return target
