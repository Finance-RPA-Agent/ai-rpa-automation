"""Readable, formula-safe Excel audit snapshots for folder reconciliation."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation, localcontext
import json
import math
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from backend.app.tools import read_data_file, validate_data
from backend.app.tools.rules import normalise_identifier, parse_amount
from product.service import result
from product.storage import timestamp, digest

INK, BLUE, MUTED = "17202F", "245DE8", "667085"
NUMBER = '#,##0.00;[Red](#,##0.00);"—"'
FIELD_LABELS = {"supplier_id": "Supplier ID", "supplier_name": "Supplier name", "order_id": "Order ID",
                "amount": "Amount", "date": "Date", "payment_status": "Payment status", "record_status": "Record status"}


def text(value):
    value = "" if value is None else str(value)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda match: f"\\u{ord(match.group()):04x}", value)[:32750]


def numeric_decimal(number):
    return float(number) if abs(number) < Decimal("10000000000000") else text(number)


def amount(value):
    try:
        return numeric_decimal(parse_amount(value))
    except (InvalidOperation, ValueError, TypeError):
        return text(value)


def difference(left, right):
    try:
        with localcontext() as context:
            context.prec = max(60, len(str(left)) + 4, len(str(right)) + 4)
            a, b = parse_amount(left), parse_amount(right)
            return numeric_decimal(a - b)
    except (InvalidOperation, ValueError, TypeError):
        return ""


def put(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    # Untrusted business text must never become an Excel formula or external link.
    if isinstance(value, str):
        cell.value = text(value)
        cell.data_type = "s"
    cell.font = Font(name="Calibri", size=11, color=INK)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def sheet_table(workbook, name, title, subtitle, columns, rows, *, numeric=()):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    put(sheet, 1, 1, title).font = Font(name="Calibri", size=22, bold=True, color=INK)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(columns), 7))
    put(sheet, 2, 1, subtitle).font = Font(name="Calibri", size=11, color=MUTED)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(len(columns), 10))
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 32
    sheet.row_dimensions[5].height = 32
    for index, (label, width) in enumerate(columns, 1):
        cell = put(sheet, 5, index, label)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number, values in enumerate(rows, 6):
        lines = max((sum(max(1, math.ceil(len(line) / max(width - 2, 1))) for line in text(value).split("\n"))
                     for value, (_, width) in zip(values, columns)), default=1)
        sheet.row_dimensions[row_number].height = min(180, max(34, lines * 14 + 8))
        for index, value in enumerate(values, 1):
            cell = put(sheet, row_number, index, value)
            if index in numeric and isinstance(value, (int, float)):
                cell.number_format = NUMBER
                cell.alignment = Alignment(horizontal="right", vertical="top")
            if name == "Reconciliation" and values[0] not in ("Reconciled",):
                if index == 1:
                    cell.fill = PatternFill("solid", fgColor="FFF2D6")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="965A00")
    if rows:
        table = Table(displayName="Batch" + name.replace(" ", ""), ref=f"A5:{get_column_letter(len(columns))}{5 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    else:
        put(sheet, 6, 1, "No outstanding items." if name == "Exceptions" else "No records to display.")
    sheet.freeze_panes = "G6" if name == "Reconciliation" else "E6" if name == "Exceptions" else "B6"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if name == "Reconciliation" else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:5"
    sheet.oddFooter.center.text = "Finance Flow | Page &P of &N"
    sheet.oddFooter.right.text = "Local reconciliation"
    return sheet


def create_batch_report(engine, view, target):
    service = engine.service
    entries = engine.entries(view["id"])
    by_id = {entry["id"]: entry for entry in entries}
    associations = engine._jobs(view["id"])
    task_details = {job["job_id"]: service.detail(job["job_id"], full=True) for job in associations}
    reconciliations, exceptions, files, decisions = [], [], [], []
    for entry in entries:
        state = entry["state"]
        if state.get("snapshot") and digest(service.path(state["snapshot"])) != state["file_sha256"]:
            raise ValueError("A captured original file changed. Its report evidence is no longer intact.")
        if state.get("canonical") and digest(service.path(state["canonical"])) != state["canonical_sha256"]:
            raise ValueError("A captured table changed. The Excel report cannot be generated from altered evidence.")
        public = next(item for item in view["tables"] if item["id"] == entry["id"])
        reference = by_id.get(entry["reference_id"])
        file_note = "; ".join(public["problems"])
        if entry["role"] == "skip":
            file_note = "Ignored"
        elif not file_note:
            file_note = "Processed" if state.get("processed") else "Validated reference" if entry["role"] == "reference" else "Not processed"
        files.append([entry["relative_path"], entry["sheet"], entry["role"], state.get("row_count", 0),
                      reference["relative_path"] + " / " + reference["sheet"] if reference else "",
                      ", ".join(state.get("currencies", [])) or "Not specified",
                      state.get("file_sha256", ""), "; ".join(f"{field} <- {column}" for field, column in state.get("mapping", {}).items()),
                      file_note])
        if entry["role"] != "source" or not state.get("canonical"):
            if entry["role"] != "skip" and public["problems"]:
                exceptions.append([entry["relative_path"], entry["sheet"], "", "", "", "Table setup", "", "", "; ".join(public["problems"]), "", "Correct the file or its table setup"])
            continue
        original = result(read_data_file(str(service.path(state["canonical"])), settings=service.settings))["records"]
        refs = defaultdict(list)
        if reference and reference["state"].get("canonical"):
            reference_rows = result(read_data_file(str(service.path(reference["state"]["canonical"])), settings=service.settings))["records"]
            for record in reference_rows:
                data = record["data"]
                refs[(normalise_identifier(data.get("supplier_id")), normalise_identifier(data.get("order_id")))].append(data)
        current, issues, task_for_row, receipts = {}, defaultdict(list), {}, {}
        for association in [job for job in associations if job["entry_id"] == entry["id"]]:
            task = task_details[association["job_id"]]
            origins = association["origins"]
            for index, origin in enumerate(origins):
                task_for_row[origin] = task
            for record in task["records"]:
                if 0 <= record["row"] - 2 < len(origins):
                    current[origins[record["row"] - 2]] = record["data"]
            for issue in task["issues"]:
                index = (issue.get("row") or 0) - 2
                if 0 <= index < len(origins):
                    issues[origins[index]].append(issue)
            if task["status"] == "waiting_confirmation":
                _, saved = service.get_job(task["id"])
                checked = result(validate_data(str(service.path(saved["working_file"])),
                                 reference_records=[items[0] for items in refs.values() if len(items) == 1], settings=service.settings))
                for issue in checked["issues"]:
                    index = (issue.get("row") or 0) - 2
                    if 0 <= index < len(origins):
                        origin = origins[index]
                        if not any(old.get("field") == issue.get("field") and old["type"] == issue["type"] for old in issues[origin]):
                            issues[origin].append(issue)
            for receipt in task["archive"]:
                try:
                    row_index = int(receipt["operation_id"].rsplit(":", 1)[1]) - 2
                    if 0 <= row_index < len(origins):
                        receipts[origins[row_index]] = receipt["id"]
                except (ValueError, IndexError):
                    pass
            for decision in task["decisions"]:
                for change in decision["changes"]:
                    index = change["source_row"] - 2
                    origin = origins[index] if 0 <= index < len(origins) else ""
                    decisions.append([entry["relative_path"], entry["sheet"], origin, task["code"], decision["revision"],
                                      FIELD_LABELS.get(change["field"], change["field"]), text(change["before"]), text(change["after"]), decision["actor"], decision["reason"], decision["created_at"]])
        external = set(state.get("cross_duplicate_indexes", []))
        for index, record in enumerate(original):
            origin = state["origin_rows"][index]
            source = record["data"]
            data = current.get(origin, source)
            task = task_for_row.get(origin)
            matches = refs[(normalise_identifier(data.get("supplier_id")), normalise_identifier(data.get("order_id")))]
            ref = matches[0] if len(matches) == 1 else {}
            row_issues = issues[origin]
            if index in external:
                status, reason = "Source correction required", "Business key is repeated in another source table."
            elif task and task["status"] == "completed":
                status, reason = "Reconciled", ""
            elif task and origin in receipts:
                status, reason = "Partial archive", task["error"] or "The task stopped after writing this receipt. Review its audit trail."
            elif task:
                status = "Needs review" if task["status"] == "waiting_confirmation" else task["status"].capitalize()
                reason = "; ".join(issue["message"] for issue in row_issues) or task["error"] or "Resume the task to revalidate the saved corrections."
            else:
                status, reason = "Not processed", "; ".join(public["problems"]) or "Start this batch to process the table."
            reconciliations.append([status, entry["relative_path"], entry["sheet"], origin, text(data.get("supplier_id")), text(data.get("order_id")),
                                    amount(source.get("amount")), amount(data.get("amount")), amount(ref.get("amount")), difference(data.get("amount"), ref.get("amount")),
                                    text(source.get("payment_status")), text(data.get("payment_status")), text(ref.get("payment_status")),
                                    text(source.get("record_status")), text(data.get("record_status")), text(ref.get("record_status")),
                                    text(source.get("supplier_name")), text(data.get("supplier_name")), text(source.get("date")), text(data.get("date")),
                                    reason, task["code"] if task else "", receipts.get(origin, "")])
            if status != "Reconciled":
                if row_issues:
                    for issue in row_issues:
                        field = issue.get("field")
                        exceptions.append([entry["relative_path"], entry["sheet"], origin, text(data.get("supplier_id")), text(data.get("order_id")),
                                           FIELD_LABELS.get(field, field or "Record"), text(data.get(field)), text(ref.get(field)), issue["message"], task["code"], "Open task to review the finding"])
                else:
                    exceptions.append([entry["relative_path"], entry["sheet"], origin, text(data.get("supplier_id")), text(data.get("order_id")),
                                       "Record", "", "", reason, task["code"] if task else "", "Open task" if task else "Correct the source or table setup"])
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 45
    summary.column_dimensions["C"].width = 26
    summary.column_dimensions["D"].width = 26
    summary.merge_cells("A1:D1")
    put(summary, 1, 1, "Reconciliation report").font = Font(name="Calibri", size=25, bold=True, color=INK)
    summary.row_dimensions[1].height = 40
    put(summary, 2, 1, view["code"]).font = Font(name="Calibri", size=13, color=BLUE, bold=True)
    metrics = [
        ("Report status", "FINAL - completed" if view["status"] == "completed" else "INTERIM - " + view["status"].replace("_", " ")),
        ("Generated at (UTC)", timestamp()), ("Started by", view["actor"]),
        ("Source tables", view["source_tables"]), ("Reference tables", view["reference_tables"]),
        ("Input source rows", view["input_rows"]), ("Completed input rows", view["matched_rows"]),
        ("Rows not completed", view["unresolved_rows"]), ("Unique archive receipts", view["archive_receipts"]),
        ("Previously created tasks reused", view["reused_tasks"]),
        ("Outstanding finding lines", len(exceptions)), ("Folder", view["folder"]),
    ]
    if view.get("bi_folder"): metrics.append(("BI system folder", view["bi_folder"]))
    for index, (label, value) in enumerate(metrics, 4):
        summary.row_dimensions[index].height = 30 if label not in ("Folder", "BI system folder") else 48
        put(summary, index, 1, label).font = Font(name="Calibri", size=11, color=MUTED)
        put(summary, index, 2, value).font = Font(name="Calibri", size=12, color=INK, bold=True)
        for column in (1, 2):
            summary.cell(index, column).border = Border(bottom=Side(style="hair", color="E2E5EB"))
        if label in ("Folder", "BI system folder"):
            summary.merge_cells(start_row=index, start_column=2, end_row=index, end_column=4)
    for column, title, destination in [(1, "Open exceptions →", "Exceptions"), (2, "Open all records →", "Reconciliation")]:
        cell = put(summary, 17, column, title)
        cell.hyperlink = f"#'{destination}'!A1"
        cell.font = Font(name="Calibri", size=12, color=BLUE, underline="single")
    notes = ["Start with Exceptions for the items that still need a decision. Reconciliation contains every captured source row.",
             "Amounts retain the supplied currency. One currency per matched table pair is assumed when no currency column is present. No FX conversion is performed.",
             "Matching uses supplier ID and order ID. Amount, payment status and record status are compared; names and dates are validated but not cross-compared.",
             "Source row refers to the original worksheet row or CSV record row, including the header. Original files and review decisions are retained in the local workspace.",
             "This is a generated audit snapshot, not an editable calculation model. Long text is limited by Excel's cell size. Large amounts outside Excel's reliable numeric precision are kept as text.",
             "New versions may create new reconciliation receipts. Identical table subsets and reference snapshots reuse existing tasks; no payment is executed."]
    for index, note in enumerate(notes, 19):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=4)
        put(summary, index, 1, note).font = Font(name="Calibri", size=11, color=MUTED)
        summary.row_dimensions[index].height = 35
    summary.freeze_panes = "B4"
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.print_options.horizontalCentered = True
    if view.get("command"):
        summary.merge_cells("A26:D26")
        put(summary, 26, 1, "Original request").font = Font(name="Calibri", size=12, bold=True, color=INK)
        summary.merge_cells("A27:D30")
        put(summary, 27, 1, view["command"]["instruction"])
        summary.row_dimensions[27].height = 60
    summary.print_area = "A1:D30" if view.get("command") else "A1:D24"
    exception_sheet = sheet_table(workbook, "Exceptions", "Items needing attention", "Only outstanding items are shown. Click a source row to open its full reconciliation record.",
                [("Source file", 32), ("Worksheet", 20), ("Source row", 12), ("Supplier ID", 20), ("Order ID", 20), ("Field", 22),
                 ("Current value", 30), ("BI value", 30), ("Finding", 65), ("Task", 16), ("Next action", 40)], exceptions)
    row_positions = {(row[1], row[2], row[3]): index + 6 for index, row in enumerate(reconciliations)}
    for index, row in enumerate(exceptions, 6):
        position = row_positions.get((row[0], row[1], row[2]))
        if position is not None:
            cell = exception_sheet.cell(index, 3)
            cell.hyperlink = f"#'Reconciliation'!A{position}"
            cell.font = Font(name="Calibri", size=11, color=BLUE, underline="single")
    sheet_table(workbook, "Reconciliation", "All source records", "Original and current values are kept side by side. Filter Result to inspect unresolved rows.",
                [("Result", 24), ("Source file", 32), ("Worksheet", 20), ("Source row", 12), ("Supplier ID", 20), ("Order ID", 20),
                 ("Source amount", 18), ("Current amount", 18), ("BI amount", 18), ("Difference", 18),
                 ("Source payment", 20), ("Current payment", 20), ("BI payment", 20),
                 ("Source record status", 22), ("Current record status", 22), ("BI record status", 22),
                 ("Source supplier name", 34), ("Current supplier name", 34), ("Source date", 18), ("Current date", 18),
                 ("Finding / reason", 65), ("Task", 16), ("Archive receipt", 26)], reconciliations, numeric=(7, 8, 9, 10))
    sheet_table(workbook, "File log", "File and table provenance", "Each worksheet is identified separately. File hashes identify the captured original bytes.",
                [("File", 38), ("Worksheet", 22), ("Role", 16), ("Rows", 12), ("Selected reference", 48), ("Currency", 18), ("Original SHA-256", 68),
                 ("Field mapping", 90), ("Notes", 80)], files)
    sheet_table(workbook, "Review history", "Recorded human decisions", "Authenticated reviewer, source row, revision and reason are retained for every field correction.",
                [("Source file", 32), ("Worksheet", 22), ("Source row", 12), ("Task", 16), ("Revision", 12), ("Field", 24),
                 ("Before", 32), ("After", 32), ("Reviewer", 24), ("Reason", 65), ("Recorded at (UTC)", 30)], decisions)
    workbook.properties.title = f"Finance Flow {view['code']} reconciliation report"
    workbook.properties.creator = "Finance Flow Local"
    workbook.save(target)
    workbook.close()
