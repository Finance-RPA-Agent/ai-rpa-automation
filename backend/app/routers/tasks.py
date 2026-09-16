from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy.orm import Session
from datetime import datetime
from openpyxl import load_workbook
import os
import shutil

from backend.app.database import SessionLocal
from backend.app.models import (
    Task as TaskModel,
    TaskLog,
    ExcelException,
    BIInventory,
    User
)


router = APIRouter(
    prefix="/tasks",
    tags=["Tasks"]
)


class TaskCreate(BaseModel):
    description: str
    user_id: int | None = None


class TaskResponse(BaseModel):
    id: int
    description: str
    status: str
    user_id: int | None
    created_at: datetime | None
    result: str | None

    class Config:
        from_attributes = True


class TaskStatusUpdate(BaseModel):
    status: str


class TaskResultUpdate(BaseModel):
    result: str


class TaskLogCreate(BaseModel):
    action: str
    status: str
    message: str | None = None


class TaskLogResponse(BaseModel):
    id: int
    task_id: int
    action: str
    status: str
    message: str | None
    created_at: datetime

    class Config:
        from_attributes = True


class ExceptionHandleRequest(BaseModel):
    action: str
    message: str | None = None
    new_value: str | int | float | None = None


def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()


def modify_excel_file(
    filename: str,
    row_number: int,
    column_name: str,
    new_value
):
    input_path = os.path.join(
        "uploads",
        filename
    )

    if not os.path.exists(input_path):
        raise HTTPException(
            status_code=404,
            detail="Original Excel file not found"
        )

    os.makedirs(
        "outputs",
        exist_ok=True
    )

    output_filename = f"processed_{filename}"

    output_path = os.path.join(
        "outputs",
        output_filename
    )

    if not os.path.exists(output_path):
        shutil.copy2(
            input_path,
            output_path
        )

    workbook = load_workbook(output_path)
    worksheet = workbook.active

    target_column = None

    for cell in worksheet[1]:
        if str(cell.value).strip() == str(column_name).strip():
            target_column = cell.column
            break

    if target_column is None:
        workbook.close()

        raise HTTPException(
            status_code=400,
            detail=f"Column '{column_name}' was not found in Excel"
        )

    worksheet.cell(
        row=row_number,
        column=target_column
    ).value = new_value

    workbook.save(output_path)
    workbook.close()

    return output_path


@router.post("/", response_model=TaskResponse)
def create_task(
    task: TaskCreate,
    db: Session = Depends(get_db)
):
    if task.user_id is not None:
        user = db.query(User).filter(
            User.id == task.user_id
        ).first()

        if user is None:
            raise HTTPException(
                status_code=404,
                detail="User not found"
            )

    new_task = TaskModel(
        description=task.description,
        status="pending",
        user_id=task.user_id
    )

    db.add(new_task)
    db.commit()
    db.refresh(new_task)

    return new_task


@router.get("/", response_model=list[TaskResponse])
def get_tasks(
    user_id: int | None = None,
    db: Session = Depends(get_db)
):
    query = db.query(TaskModel)

    if user_id is not None:
        query = query.filter(
            TaskModel.user_id == user_id
        )

    return query.all()


@router.get("/{task_id}", response_model=TaskResponse)
def get_task(
    task_id: int,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    return task


@router.patch("/{task_id}/status", response_model=TaskResponse)
def update_task_status(
    task_id: int,
    task_update: TaskStatusUpdate,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    task.status = task_update.status

    db.commit()
    db.refresh(task)

    return task


@router.patch("/{task_id}/result", response_model=TaskResponse)
def update_task_result(
    task_id: int,
    task_update: TaskResultUpdate,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    task.result = task_update.result

    db.commit()
    db.refresh(task)

    return task


@router.post("/{task_id}/logs", response_model=TaskLogResponse)
def create_task_log(
    task_id: int,
    log: TaskLogCreate,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    new_log = TaskLog(
        task_id=task_id,
        action=log.action,
        status=log.status,
        message=log.message
    )

    db.add(new_log)
    db.commit()
    db.refresh(new_log)

    return new_log


@router.get("/{task_id}/logs", response_model=list[TaskLogResponse])
def get_task_logs(
    task_id: int,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    return db.query(TaskLog).filter(
        TaskLog.task_id == task_id
    ).all()


@router.post("/{task_id}/exceptions/{exception_id}/handle")
def handle_exception(
    task_id: int,
    exception_id: int,
    request: ExceptionHandleRequest,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    exception = db.query(ExcelException).filter(
        ExcelException.id == exception_id,
        ExcelException.task_id == task_id
    ).first()

    if exception is None:
        raise HTTPException(
            status_code=404,
            detail="Exception not found"
        )

    allowed_actions = [
        "approve",
        "reject",
        "modify"
    ]

    if request.action not in allowed_actions:
        raise HTTPException(
            status_code=400,
            detail="Action must be approve, reject, or modify"
        )

    if request.action == "modify" and request.new_value is None:
        raise HTTPException(
            status_code=400,
            detail="new_value is required when action is modify"
        )

    output_file = None

    if request.action == "modify":
        output_file = modify_excel_file(
            filename=exception.filename,
            row_number=exception.row,
            column_name=exception.column,
            new_value=request.new_value
        )

        exception.status = "modified"
        exception.resolution = (
            request.message
            or f"Value changed to {request.new_value}"
        )

    elif request.action == "approve":
        exception.status = "approved"
        exception.resolution = request.message

    elif request.action == "reject":
        exception.status = "rejected"
        exception.resolution = request.message

    new_log = TaskLog(
        task_id=task_id,
        action=f"Exception {exception_id} {request.action}",
        status="handled",
        message=(
            request.message
            or f"Exception {exception_id} was {request.action}."
        )
    )

    db.add(new_log)
    db.commit()

    db.refresh(exception)
    db.refresh(new_log)

    return {
        "success": True,
        "task_id": task_id,
        "exception_id": exception_id,
        "action": request.action,
        "status": exception.status,
        "resolution": exception.resolution,
        "new_value": (
            request.new_value
            if request.action == "modify"
            else None
        ),
        "output_file": output_file,
        "log_id": new_log.id
    }


@router.post("/{task_id}/bi-check")
def bi_check(
    task_id: int,
    material_code: str | None = None,
    warehouse: str | None = None,
    db: Session = Depends(get_db)
):
    task = db.query(TaskModel).filter(
        TaskModel.id == task_id
    ).first()

    if task is None:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    query = db.query(BIInventory)

    if material_code:
        query = query.filter(
            BIInventory.material_code == material_code
        )

    if warehouse:
        query = query.filter(
            BIInventory.warehouse == warehouse
        )

    inventory = query.all()

    log_message = (
        f"BI inventory query completed. "
        f"{len(inventory)} records found."
    )

    new_log = TaskLog(
        task_id=task_id,
        action="BI Inventory Query",
        status="completed",
        message=log_message
    )

    db.add(new_log)
    db.commit()

    return {
        "success": True,
        "task_id": task_id,
        "material_code": material_code,
        "warehouse": warehouse,
        "count": len(inventory),
        "data": [
            {
                "id": item.id,
                "material_code": item.material_code,
                "warehouse": item.warehouse,
                "quantity": item.quantity
            }
            for item in inventory
        ]
    }