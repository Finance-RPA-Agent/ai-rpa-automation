"""Regression tests for the first independent acceptance round (F1-F7).

Each test mirrors an independent finding and adds neighbouring boundary
cases.  The independent acceptance files under
``backend-review/acceptance-round1`` are never modified.
"""

from __future__ import annotations

import csv

import pytest

from backend.app.models import OaReconciliationRecord, Task
from backend.app.services.demo_workflow import (
    DemoFlowError,
    run_conflict_flow,
    run_normal_flow,
)
from backend.app.tools import read_data_file, submit_to_oa, validate_data

FIELDS = [
    "supplier_id",
    "supplier_name",
    "order_id",
    "amount",
    "date",
    "payment_status",
    "record_status",
]
ROW = [
    "SUP-001",
    "ABC Ltd",
    "ORD-1001",
    "5000.00",
    "2026-09-18",
    "pending",
    "active",
]
REFERENCE = dict(zip(FIELDS, ROW))


def _write(settings, rows, name="acceptance.csv", encoding="utf-8"):
    path = settings.upload_dir / name
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding=encoding, newline="") as handle:
        csv.writer(handle).writerows(rows)
    return path


def _issue_types(result):
    return [issue["type"] for issue in result["issues"]]


# ----------------------------------------------------------------- F1
def test_duplicate_required_header_is_reported(temp_settings):
    path = _write(temp_settings, [FIELDS + ["amount"], ROW + ["5200.00"]])

    response = validate_data(
        str(path),
        reference_records=[REFERENCE],
        settings=temp_settings,
    )

    assert response["status"] == "success"
    assert "duplicate_header" in _issue_types(response["result"])


def test_duplicate_amount_header_cannot_complete_normal_flow(
    db_session, temp_settings
):
    path = _write(temp_settings, [FIELDS + ["amount"], ROW + ["5200.00"]])

    try:
        run_normal_flow(db_session, file_path=path, settings=temp_settings)
    except DemoFlowError:
        pass
    else:  # pragma: no cover - the flow must not report success
        raise AssertionError("normal flow accepted ambiguous headers")

    task = db_session.query(Task).order_by(Task.id.desc()).first()
    assert task.status != "completed"
    assert db_session.query(OaReconciliationRecord).count() == 0


def test_blank_and_whitespace_headers_are_reported(make_csv, temp_settings):
    rows = [
        ["supplier_id", "   ", "amount"],
        ["SUP-001", "x", "10.00"],
    ]
    path = make_csv(rows)

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    types = {
        issue["type"] for issue in response["result"]["column_issues"]
    }
    assert "empty_header" in types
    assert len(set(response["result"]["columns"])) == 3


# ----------------------------------------------------------------- F2
def test_header_collision_preserves_every_cell(make_csv, temp_settings):
    rows = [
        ["amount", "amount", "amount__2"],
        ["5000", "5200", "5300"],
    ]
    path = make_csv(rows)

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert len(set(result["columns"])) == 3
    data = result["records"][0]["data"]
    assert len(data) == 3
    assert [data[column] for column in result["columns"]] == [
        "5000",
        "5200",
        "5300",
    ]


def test_generated_header_name_collision_is_resolved(make_csv, temp_settings):
    rows = [
        [None, "column_1", "value"],
        ["a", "b", "c"],
    ]
    path = make_csv(rows)

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert len(set(result["columns"])) == 3
    assert len(result["records"][0]["data"]) == 3


def test_repeated_generated_names_stay_unique(make_csv, temp_settings):
    rows = [
        [None, None, None],
        ["a", "b", "c"],
    ]
    path = make_csv(rows)

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    columns = response["result"]["columns"]
    assert len(set(columns)) == 3


# ----------------------------------------------------------------- F3
@pytest.mark.parametrize(
    "value",
    [
        "50,00",
        "5,,000",
        ",5000",
        "1,,000.00",
        "12,34,567",
        "1,2345.00",
        "1234,567",
        "1,250.",
        ".50",
        "5 000",
        "1,250.123",
        "abc",
    ],
)
def test_malformed_amounts_are_rejected(make_csv, temp_settings, value):
    row = list(ROW)
    row[3] = value
    path = make_csv([FIELDS, row])

    response = validate_data(str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert "invalid_amount" in _issue_types(response["result"])


@pytest.mark.parametrize(
    "value",
    ["1,250.50", "1,250", "0.5", "1234567.89", "0", "0.00", "1,000,000.01"],
)
def test_legal_amounts_still_parse(make_csv, temp_settings, value):
    row = list(ROW)
    row[3] = value
    path = make_csv([FIELDS, row])

    response = validate_data(str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert "invalid_amount" not in _issue_types(response["result"])


def test_oa_amount_check_uses_same_strict_rule(db_session):
    task = Task(description="strict amount", status="running")
    db_session.add(task)
    db_session.commit()

    response = submit_to_oa(
        task_id=task.id,
        record_type="reconciliation_result",
        data=dict(
            REFERENCE,
            amount="50,00",
            record_status="matched",
        ),
        operation_id="strict-amount-op",
        db=db_session,
    )

    assert response["status"] == "error"
    assert response["error_code"] == "INVALID_DATA"


# ----------------------------------------------------------------- F4
@pytest.mark.parametrize("field", ["payment_status", "record_status"])
def test_missing_reference_status_is_reported(
    make_csv, temp_settings, field
):
    reference = {k: v for k, v in REFERENCE.items() if k != field}
    path = make_csv([FIELDS, ROW])

    response = validate_data(
        str(path),
        reference_records=[reference],
        settings=temp_settings,
    )

    assert response["status"] == "success"
    result = response["result"]
    assert "missing_reference_status" in _issue_types(result)
    assert result["invalid_rows"] == 1


@pytest.mark.parametrize("field", ["payment_status", "record_status"])
def test_blank_reference_status_is_reported(
    make_csv, temp_settings, field
):
    reference = dict(REFERENCE)
    reference[field] = "   "
    path = make_csv([FIELDS, ROW])

    response = validate_data(
        str(path),
        reference_records=[reference],
        settings=temp_settings,
    )

    assert "missing_reference_status" in _issue_types(
        response["result"]
    )


def test_non_scalar_reference_status_is_reported(make_csv, temp_settings):
    reference = dict(REFERENCE)
    reference["payment_status"] = ["pending"]
    path = make_csv([FIELDS, ROW])

    response = validate_data(
        str(path),
        reference_records=[reference],
        settings=temp_settings,
    )

    assert "invalid_reference_status" in _issue_types(response["result"])


def test_complete_reference_statuses_are_not_over_flagged(
    make_csv, temp_settings
):
    path = make_csv([FIELDS, ROW])

    response = validate_data(
        str(path),
        reference_records=[REFERENCE],
        settings=temp_settings,
    )

    assert response["result"]["issues"] == []


# ----------------------------------------------------------------- F5
@pytest.mark.parametrize(
    "bad_field,bad_value",
    [("amount", "bad"), ("date", "2026/09/18")],
)
def test_invalid_only_input_does_not_complete_conflict_flow(
    db_session, temp_settings, bad_field, bad_value
):
    row = list(ROW)
    row[FIELDS.index(bad_field)] = bad_value
    path = _write(temp_settings, [FIELDS, row], name="invalid-only.csv")

    try:
        run_conflict_flow(db_session, file_path=path, settings=temp_settings)
    except DemoFlowError:
        pass
    else:  # pragma: no cover
        raise AssertionError("conflict flow treated format errors as done")

    task = db_session.query(Task).order_by(Task.id.desc()).first()
    assert task.status == "waiting_confirmation"
    assert db_session.query(OaReconciliationRecord).count() == 0

    import json

    summary = json.loads(task.result)
    assert summary["unarchived_issues"]
    assert summary["normal_result_recorded"] is False


def test_mixed_conflict_and_format_errors_hold_for_review(
    db_session, temp_settings
):
    rows = [
        FIELDS,
        ["SUP-001", "ABC Ltd", "ORD-1001", "5200.00", "2026-09-18", "pending", "active"],
        ["SUP-002", "Bright Trading", "ORD-2001", "bad", "2026-09-20", "pending", "active"],
    ]
    path = _write(temp_settings, rows, name="mixed.csv")

    with pytest.raises(DemoFlowError):
        run_conflict_flow(db_session, file_path=path, settings=temp_settings)

    task = db_session.query(Task).order_by(Task.id.desc()).first()
    assert task.status == "waiting_confirmation"

    import json

    summary = json.loads(task.result)
    assert summary["conflicts"]
    assert summary["unarchived_issues"]
    assert summary["normal_result_recorded"] is False

    records = db_session.query(OaReconciliationRecord).all()
    assert len(records) == 1
    assert records[0].record_type == "reconciliation_exception"
    assert records[0].review_status == "pending_review"


# ----------------------------------------------------------------- F6
@pytest.mark.parametrize("identifier", ["NA", "N/A", "NULL", "None", "nan"])
def test_xlsx_text_placeholders_are_preserved(
    temp_settings, identifier
):
    from openpyxl import Workbook

    path = temp_settings.upload_dir / f"identifier-{identifier.replace('/', '_')}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    row = list(ROW)
    row[0] = identifier
    workbook = Workbook()
    workbook.active.append(FIELDS)
    workbook.active.append(row)
    workbook.save(path)
    workbook.close()

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert (
        response["result"]["records"][0]["data"]["supplier_id"]
        == identifier
    )


def test_xlsx_real_empty_cell_stays_missing(temp_settings):
    from openpyxl import Workbook

    path = temp_settings.upload_dir / "empty-cell.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    row = list(ROW)
    row[1] = None
    workbook = Workbook()
    workbook.active.append(FIELDS)
    workbook.active.append(row)
    workbook.save(path)
    workbook.close()

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert response["result"]["records"][0]["data"]["supplier_name"] is None
    assert response["result"]["records"][0]["row"] == 2


# ----------------------------------------------------------------- F7
def test_csv_blank_line_does_not_shift_issue_location(
    make_csv, temp_settings
):
    row = list(ROW)
    row[3] = "bad"
    path = make_csv([FIELDS, [], row], name="blank-line.csv")

    response = validate_data(str(path), settings=temp_settings)

    assert response["status"] == "success"
    issue = next(
        issue
        for issue in response["result"]["issues"]
        if issue["type"] == "invalid_amount"
    )
    assert issue["row"] == 3


def test_multiple_csv_blank_lines_keep_line_numbers(
    make_csv, temp_settings
):
    row = list(ROW)
    row[3] = "bad"
    path = make_csv(
        [FIELDS, [], [], row], name="blank-lines.csv"
    )

    response = validate_data(str(path), settings=temp_settings)

    issue = next(
        issue
        for issue in response["result"]["issues"]
        if issue["type"] == "invalid_amount"
    )
    assert issue["row"] == 4


def test_quoted_newline_keeps_physical_start_lines(
    make_csv, temp_settings
):
    first = list(ROW)
    first[1] = "ABC\nLtd"
    first[3] = "bad"
    second = [
        "SUP-002",
        "Bright Trading",
        "ORD-2001",
        "",
        "2026-09-20",
        "pending",
        "active",
    ]
    path = make_csv(
        [FIELDS, first, second], name="quoted-newline.csv"
    )

    response = validate_data(str(path), settings=temp_settings)

    amount_issue = next(
        issue
        for issue in response["result"]["issues"]
        if issue["type"] == "invalid_amount"
    )
    missing_issue = next(
        issue
        for issue in response["result"]["issues"]
        if issue["type"] == "missing_value"
    )
    assert amount_issue["row"] == 2
    assert missing_issue["row"] == 4


def test_read_csv_reports_physical_row_numbers(
    make_csv, temp_settings
):
    path = make_csv([FIELDS, [], ROW], name="rows.csv")

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert response["result"]["row_count"] == 1
    assert response["result"]["records"][0]["row"] == 3


def test_quoted_comma_field_keeps_single_record(make_csv, temp_settings):
    first = list(ROW)
    first[1] = "ABC, Ltd"
    path = make_csv(
        [FIELDS, first, ROW], name="quoted-comma.csv"
    )

    response = read_data_file(str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert result["row_count"] == 2
    assert result["records"][0]["data"]["supplier_name"] == "ABC, Ltd"
    assert result["records"][0]["row"] == 2
    assert result["records"][1]["row"] == 3
