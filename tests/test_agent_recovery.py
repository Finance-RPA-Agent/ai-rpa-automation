"""Regression for invalid tool arguments, phantom count columns and retained retries."""
import csv
import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import httpx
import pytest

from product.api import create_app
from product.data_contracts import action_schema
from product.data_tools import DataTools
from product.model import OllamaProvider

REQUEST = "Combine the orders in the Input folder, match them against the BI system folder by supplier ID and order ID, calculate amount differences, summarize order counts and both amounts by payment status, create a difference bar chart, and export the details, summary, and chart to an Excel report with a brief explanation of the findings."
FIELDS = ["supplier_id", "supplier_name", "order_id", "amount", "date", "payment_status", "record_status"]
ROWS = [["001", "Acme", "0001", "100.00", "2026-10-01", "pending", "active"],
        ["002", "Cedar", "0002", "200.00", "2026-10-02", "paid", "active"],
        ["003", "Delta", "0003", "300.00", "2026-10-03", "pending", "active"]]
GOALS = ["list_files", "inspect_table", "concatenate", "join_tables", "derive_column", "aggregate", "create_chart", "export_report"]


def call(tool, **arguments): return {"action": "tool", "tool": tool, "arguments": arguments, "message": ""}
def done(): return {"action": "final", "tool": None, "arguments": {}, "message": "The actual results and Excel report are available."}


class Model:
    def __init__(self): self.actions, self.goals = [], list(GOALS)
    def data_plan(self, context, schema): return {"supported": True, "required_tools": self.goals, "message": "Complete the requested outputs."}
    def data_action(self, context, schema): return self.actions.pop(0) if self.actions else {"action": "ask", "message": "Which operation should I use?"}


@pytest.fixture
def local(tmp_path):
    source, bi = tmp_path / "input", tmp_path / "bi"
    changed = [list(row) for row in ROWS]; changed[0][3] = "125.00"
    for path, rows in ((source / "part1.csv", changed[:2]), (source / "part2.csv", changed[2:]), (bi / "reference.csv", ROWS)):
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream); writer.writerow(FIELDS); writer.writerows(rows)
    model = Model(); app = create_app(tmp_path / "workspace", model=model)
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Requested-With": "FinanceFlow"}) as client:
        client.post("/api/setup", json={"username": "Recovery Owner", "password": "123456"})
        chat = client.post("/api/assistant/chats", json={"folder": str(source), "bi_folder": str(bi)}).json()
        yield client, app.state.service, model, chat


def run(local, actions, retry=None, instruction=REQUEST):
    client, service, model, chat = local; model.actions = list(actions)
    body = {"instruction": instruction}
    if retry is not None: body["retry_turn_id"] = retry
    response = client.post(f"/api/assistant/chats/{chat['id']}/turns", json=body)
    assert response.status_code == 202, response.text
    service.future.result(timeout=60)
    return client.get(f"/api/assistant/chats/{chat['id']}").json()


def test_generation_schema_binds_each_tool_and_aggregate_input_columns():
    context = {"files": [{"file_id": "f1"}], "frames": [{"id": "t7", "columns": ["order_id", "amount_left", "payment_status_left"]}]}
    schema = action_schema(context)
    branches = [branch for branch in schema["anyOf"] if branch["properties"]["action"]["enum"] == ["tool"]]
    listing = next(branch for branch in branches if branch["properties"]["tool"]["enum"] == ["list_files"])
    assert listing["properties"]["arguments"]["properties"] == {}
    assert listing["properties"]["arguments"]["additionalProperties"] is False
    aggregate = next(branch for branch in branches if branch["properties"]["tool"]["enum"] == ["aggregate"])
    args = aggregate["properties"]["arguments"]
    assert args["properties"]["frame_id"]["enum"] == ["t7"]
    metric = args["properties"]["metrics"]["items"]
    assert metric["additionalProperties"] is False
    assert "order_id" in metric["properties"]["column"]["enum"]
    assert "order_count" not in metric["properties"]["column"]["enum"]
    assert "message" not in args["properties"]
    assert "$ref" not in json.dumps(schema)
    assert "enum" not in metric["properties"]["output"]


def test_schema_uses_full_frame_columns_without_exposing_private_catalog_to_model():
    requests = []
    def handle(request):
        if request.url.path == "/api/tags": return httpx.Response(200, json={"models": [{"name": "qwen3.5:4b"}]})
        body = json.loads(request.content); requests.append(body)
        assert "_frame_catalog" not in body["messages"][-1]["content"]
        assert body["format"].get("anyOf") and body["options"]["num_ctx"] == 8192
        return httpx.Response(200, json={"done": True, "message": {"content": json.dumps(call("aggregate", frame_id="t7", group_by=[], metrics=[{"column": "actual_26th_column", "op": "count", "output": "order_count"}]))}})
    context = {"files": [{"file_id": "f1"}], "frames": [{"id": "t7", "columns": ["short_preview"]}],
               "_frame_catalog": [{"id": "t7", "columns": ["short_preview", "actual_26th_column"]}]}
    result = OllamaProvider(transport=httpx.MockTransport(handle)).data_action(context, {})
    assert result["arguments"]["metrics"][0]["column"] == "actual_26th_column"
    assert "actual_26th_column" in json.dumps(requests[0]["format"])


def test_screenshot_sequence_recovers_after_three_separated_errors(local):
    selected = ["supplier_id", "order_id", "amount_left", "payment_status_left", "amount_right", "payment_status_right"]
    metrics = [{"column": "order_id", "op": "count", "output": "order_count"},
               {"column": "amount_left", "op": "sum", "output": "input_total"},
               {"column": "amount_right", "op": "sum", "output": "bi_total"},
               {"column": "amount_difference", "op": "sum", "output": "difference_total"}]
    actions = [call("list_files", file_id="f1"), call("list_files"),
        call("inspect_table", file_id="f1"), call("inspect_table", file_id="f2"),
        call("concatenate", frame_ids=["t1", "t2"]), call("inspect_table", file_id="f3"),
        call("join_tables", left_id="t3", right_id="t4", left_on=["supplier_id", "order_id"], right_on=["supplier_id", "order_id"]),
        call("select_columns", frame_id="t5", columns=selected),
        call("derive_column", frame_id="t6", column="amount_difference", operation="subtract", left="amount_left", right="amount_right"),
        call("aggregate", frame_id="t7", group_by=["payment_status_left"], metrics=metrics, message=""),
        call("aggregate", frame_id="t7", group_by=["payment_status_left"], metrics=[{"column": "order_count", "op": "count", "output": "order_count"}]),
        call("aggregate", frame_id="t7", group_by=["payment_status_left"], metrics=metrics),
        call("create_chart", frame_id="t8", x="payment_status_left", y="difference_total", kind="bar"),
        call("export_report", frame_ids=["t7", "t8"]), done()]
    result = run(local, actions); turn = result["turns"][-1]
    assert turn["status"] == "completed" and turn["error"] is None
    assert sum(step["status"] == "error" for step in turn["steps"]) == 3
    assert turn["action_budget"] > 12
    workbook = load_workbook(io.BytesIO(local[0].get(result["artifacts"][-1]["url"]).content))
    assert workbook["Result 1"].max_row == 4
    assert workbook["Result 2"].max_row == 3 and len(workbook["Chart 1"]._charts) == 1


def test_retry_preserves_successful_steps_and_does_not_read_again(local, monkeypatch):
    client, service, model, chat = local
    model.goals = ["inspect_table", "aggregate", "create_chart", "export_report"]
    request = "Summarize the input amounts by payment status, create a bar chart and export the details and summary."
    first = run(local, [call("inspect_table", file_id="f1"), call("unknown"), call("unknown"), call("unknown")], instruction=request)
    failed = first["turns"][-1]; assert failed["status"] == "failed"
    state = json.loads(service.assistant.chat(chat["id"])["state_json"])
    first_hash = state["frames"]["t1"]["sha256"]
    def forbid_read(*args, **kwargs): raise AssertionError("Retry must reuse the saved frame.")
    monkeypatch.setattr(DataTools, "inspect_table", forbid_read)
    result = run(local, [call("aggregate", frame_id="t1", group_by=["payment_status"], metrics=[{"column": "amount", "op": "sum", "output": "Total"}]),
        call("create_chart", frame_id="t2", x="payment_status", y="Total"), call("export_report", frame_ids=["t1", "t2"]), done()], retry=failed["id"], instruction=request)
    retried = result["turns"][-1]
    assert retried["status"] == "completed" and retried["retry_of"] == failed["id"]
    assert retried["steps"][0]["tool"] == "inspect_table" and retried["steps"][0]["reused"] is True
    state = json.loads(service.assistant.chat(chat["id"])["state_json"])
    assert state["frames"]["t1"]["sha256"] == first_hash
    assert result["turns"][0]["status"] == "failed"


def test_retry_cannot_change_request_or_cross_conversations(local):
    client, _, model, chat = local; model.goals = []
    failed = run(local, [call("unknown")] * 3)["turns"][-1]
    endpoint = f"/api/assistant/chats/{chat['id']}/turns"
    assert client.post(endpoint, json={"instruction": "Different request", "retry_turn_id": failed["id"]}).status_code == 422
    assert client.post(endpoint, json={"instruction": REQUEST, "retry_turn_id": 99999}).status_code == 404


def test_difference_chart_schema_cannot_choose_order_counts():
    frames = [{"id": "detail", "columns": ["payment_status", "difference"], "rows": 100, "difference_columns": ["difference"]},
              {"id": "summary", "columns": ["payment_status", "count", "delta_total"], "rows": 2, "difference_columns": ["delta_total"]}]
    schema = action_schema({"files": [], "frames": frames, "difference_chart_required": True})
    charts = [branch for branch in schema["anyOf"] if branch["properties"].get("tool", {}).get("enum") == ["create_chart"]]
    assert len(charts) == 1
    assert charts[0]["properties"]["arguments"]["properties"]["frame_id"]["enum"] == ["summary"]
    assert charts[0]["properties"]["arguments"]["properties"]["y"]["enum"] == ["delta_total"]


def test_runtime_rejects_count_chart_and_grounds_summary_in_actual_values(local):
    local[2].goals = ["inspect_table", "derive_column", "aggregate", "create_chart", "export_report"]
    result = run(local, [call("inspect_table", file_id="f1"),
        call("derive_column", frame_id="t1", column="delta", operation="subtract", left="amount", factor="0"),
        call("aggregate", frame_id="t2", group_by=["payment_status"], metrics=[{"column": "order_id", "op": "count", "output": "Count"}, {"column": "delta", "op": "sum", "output": "Difference"}]),
        call("create_chart", frame_id="t3", x="payment_status", y="Count"),
        call("create_chart", frame_id="t3", x="payment_status", y="Difference"),
        call("export_report", frame_ids=["t2", "t3"]),
        {"action": "final", "tool": None, "arguments": {}, "message": "The average difference is 999999."}],
        instruction="Calculate differences against zero and create a difference bar chart by payment status with an Excel report.")
    last = result["turns"][-1]
    assert last["status"] == "completed"
    assert any("not order counts" in step.get("error", "") for step in last["steps"])
    assert "999999" not in last["message"] and "999999" in last["model_message"]
    assert result["charts"][-1]["y"] == "Difference"


def test_retry_recovers_difference_lineage_from_old_saved_steps(local):
    client, service, model, chat = local
    model.goals = ["inspect_table", "derive_column", "aggregate", "create_chart", "export_report"]
    request = "Calculate an amount difference against zero, summarize it and create a difference bar chart with an Excel report."
    first = run(local, [call("inspect_table", file_id="f1"), call("derive_column", frame_id="t1", column="delta", operation="subtract", left="amount", factor="0"), call("unknown"), call("unknown"), call("unknown")], instruction=request)
    assert first["turns"][-1]["status"] == "failed"
    with service.store.connect() as db:
        state = json.loads(service.assistant.chat(chat["id"])["state_json"])
        for info in state["frames"].values(): info.pop("difference_columns", None)
        db.execute("UPDATE data_chats SET state_json=? WHERE id=?", (json.dumps(state), chat["id"]))
    result = run(local, [call("aggregate", frame_id="t2", group_by=["payment_status"], metrics=[{"column": "delta", "op": "sum", "output": "Difference"}]),
        call("create_chart", frame_id="t3", x="payment_status", y="Difference"), call("export_report", frame_ids=["t2", "t3"]), done()], retry=first["turns"][-1]["id"], instruction=request)
    assert result["turns"][-1]["status"] == "completed"
    assert "Computed summary:" in result["turns"][-1]["message"]
