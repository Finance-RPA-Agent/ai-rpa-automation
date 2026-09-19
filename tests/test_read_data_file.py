"""Tests for the read_data_file tool."""

from __future__ import annotations

from pathlib import Path

from backend.app.tools import read_data_file

CORE_FIELDS = [
    "supplier_id",
    "supplier_name",
    "order_id",
    "amount",
    "date",
    "payment_status",
    "record_status",
]


def _normal_rows():
    return [
        CORE_FIELDS,
        [
            "SUP-001",
            "ABC Ltd",
            "ORD-1001",
            5000.0,
            "2026-09-18",
            "pending",
            "active",
        ],
    ]


def test_read_xlsx_success(make_xlsx, temp_settings):
    path = make_xlsx(_normal_rows())

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert result["file_type"] == "xlsx"
    assert result["row_count"] == 1
    assert result["columns"] == CORE_FIELDS
    assert result["header_row"] == 1
    assert result["first_data_row"] == 2
    assert result["records"][0]["row"] == 2
    assert result["records"][0]["data"]["supplier_id"] == "SUP-001"
    assert result["records"][0]["data"]["order_id"] == "ORD-1001"
    assert result["column_issues"] == []
    assert Path(result["file_path"]).is_absolute()


def test_read_csv_utf8(make_csv, temp_settings):
    path = make_csv(_normal_rows(), name="orders.csv")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert result["file_type"] == "csv"
    assert result["sheet_name"] is None
    assert result["row_count"] == 1
    assert result["records"][0]["data"]["amount"] in ("5000.0", "5000")


def test_read_csv_with_bom(make_csv, temp_settings):
    path = make_csv(_normal_rows(), name="bom.csv", encoding="utf-8-sig")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert response["result"]["columns"][0] == "supplier_id"
    assert (
        response["result"]["records"][0]["data"]["supplier_id"] == "SUP-001"
    )


def test_leading_zero_identifier_is_preserved(make_csv, temp_settings):
    rows = [
        CORE_FIELDS,
        [
            "00123",
            "Zero Ltd",
            "ORD-0007",
            "12.50",
            "2026-09-18",
            "pending",
            "active",
        ],
    ]
    path = make_csv(rows)

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    data = response["result"]["records"][0]["data"]
    assert data["supplier_id"] == "00123"
    assert data["order_id"] == "ORD-0007"


def test_leading_zero_identifier_is_preserved_in_xlsx(make_xlsx, temp_settings):
    rows = [
        CORE_FIELDS,
        [
            "00123",
            "Zero Ltd",
            "ORD-0007",
            12.5,
            "2026-09-18",
            "pending",
            "active",
        ],
    ]
    path = make_xlsx(rows)

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert (
        response["result"]["records"][0]["data"]["supplier_id"] == "00123"
    )


def test_sheet_name_selection(make_xlsx, temp_settings):
    path = make_xlsx(_normal_rows(), name="multi.xlsx", sheet_name="First")
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    second = workbook.create_sheet("Second")
    second.append(["supplier_id"])
    second.append(["SUP-999"])
    workbook.save(path)
    workbook.close()

    response = read_data_file(
        file_path=str(path),
        sheet_name="Second",
        settings=temp_settings,
    )

    assert response["status"] == "success"
    assert response["result"]["sheet_name"] == "Second"
    assert response["result"]["columns"] == ["supplier_id"]


def test_default_sheet_is_first_sheet(make_xlsx, temp_settings):
    path = make_xlsx(_normal_rows(), name="multi.xlsx", sheet_name="First")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert response["result"]["sheet_name"] == "First"


def test_unknown_sheet_reports_sheet_not_found(make_xlsx, temp_settings):
    path = make_xlsx(_normal_rows())

    response = read_data_file(
        file_path=str(path),
        sheet_name="Missing",
        settings=temp_settings,
    )

    assert response["status"] == "error"
    assert response["error_code"] == "SHEET_NOT_FOUND"
    assert "Missing" in response["message"]


def test_file_not_found(temp_settings):
    response = read_data_file(
        file_path=str(temp_settings.upload_dir / "nope.xlsx"),
        settings=temp_settings,
    )

    assert response["status"] == "error"
    assert response["error_code"] == "FILE_NOT_FOUND"


def test_unsupported_extension(temp_settings):
    path = temp_settings.upload_dir / "data.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("hello", encoding="utf-8")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "error"
    assert response["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_legacy_xls_is_not_claimed(temp_settings):
    path = temp_settings.upload_dir / "legacy.xls"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"\xd0\xcf\x11\xe0fake-old-excel")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "error"
    assert response["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_corrupt_xlsx_reports_read_error(temp_settings):
    path = temp_settings.upload_dir / "broken.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"this is not a real xlsx file")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "error"
    assert response["error_code"] == "FILE_READ_ERROR"


def test_empty_csv_reports_empty_file(temp_settings):
    path = temp_settings.upload_dir / "empty.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("", encoding="utf-8")

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "error"
    assert response["error_code"] == "EMPTY_FILE"


def test_header_only_csv_is_success_with_zero_rows(make_csv, temp_settings):
    path = make_csv([CORE_FIELDS])

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    assert response["result"]["row_count"] == 0
    assert response["result"]["records"] == []
    assert response["result"]["columns"] == CORE_FIELDS


def test_duplicate_headers_are_reported(make_xlsx, temp_settings):
    rows = [
        ["supplier_id", "order_id", "amount", "amount"],
        ["SUP-001", "ORD-1001", 10, 20],
    ]
    path = make_xlsx(rows)

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    result = response["result"]
    assert result["row_count"] == 1
    types = [issue["type"] for issue in result["column_issues"]]
    assert "duplicate_header" in types
    record = result["records"][0]["data"]
    assert record["amount"] == 10
    assert record["amount__2"] == 20


def test_empty_headers_are_reported(make_xlsx, temp_settings):
    rows = [
        ["supplier_id", None, "amount"],
        ["SUP-001", "ORD-1001", 10],
    ]
    path = make_xlsx(rows)

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    types = [issue["type"] for issue in response["result"]["column_issues"]]
    assert "empty_header" in types
    assert len(response["result"]["columns"]) == 3


def test_path_traversal_is_rejected(temp_settings, tmp_path):
    secret = tmp_path.parent / "outside-secret.csv"
    secret.write_text("a,b\n1,2\n", encoding="utf-8")

    response = read_data_file(
        file_path="../outside-secret.csv",
        settings=temp_settings,
    )

    assert response["status"] == "error"
    assert response["error_code"] in ("PATH_NOT_ALLOWED", "FILE_NOT_FOUND")


def test_absolute_path_outside_allowed_roots_is_rejected(
    temp_settings, tmp_path
):
    outside = tmp_path.parent / "system-file.csv"
    outside.write_text("a,b\n1,2\n", encoding="utf-8")

    response = read_data_file(
        file_path=str(outside),
        settings=temp_settings,
    )

    assert response["status"] == "error"
    assert response["error_code"] == "PATH_NOT_ALLOWED"


def test_relative_path_resolves_inside_upload_dir(temp_settings, make_csv):
    source = make_csv(_normal_rows(), name="orders.csv")
    temp_settings.upload_dir.mkdir(parents=True, exist_ok=True)
    target = temp_settings.upload_dir / "orders.csv"
    target.write_bytes(source.read_bytes())

    response = read_data_file(
        file_path="orders.csv",
        settings=temp_settings,
    )

    assert response["status"] == "success"
    assert response["result"]["row_count"] == 1


def test_missing_file_path_argument(temp_settings):
    response = read_data_file(settings=temp_settings)

    assert response["status"] == "error"
    assert response["error_code"] == "INVALID_ARGUMENTS"


def test_dates_and_numbers_are_json_safe(make_xlsx, temp_settings):
    import json
    from datetime import datetime

    rows = [
        CORE_FIELDS,
        [
            "SUP-001",
            "ABC Ltd",
            "ORD-1001",
            5000.5,
            datetime(2026, 9, 18),
            "pending",
            "active",
        ],
    ]
    path = make_xlsx(rows)

    response = read_data_file(file_path=str(path), settings=temp_settings)

    assert response["status"] == "success"
    encoded = json.dumps(response)
    assert "2026-09-18" in encoded
    data = response["result"]["records"][0]["data"]
    assert data["date"] == "2026-09-18"
    assert float(data["amount"]) == 5000.5
