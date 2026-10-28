"""Separate processing and BI directories across tasks, snapshots and reports."""
import csv
import io
import json
import sqlite3

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from product.api import create_app
from product.data_tools import DataTools
from product.service import FIELDS
from product.storage import Store
from start_product import restore_backup

ROWS = [["001", "Acme", "0001", "100.00", "2026-10-01", "pending", "active"],
        ["002", "Cedar", "0002", "200.00", "2026-10-02", "paid", "active"]]


class Model:
    def interpret_command(self, instruction, schema):
        return {"intent": "reconcile_folder", "folders": [], "message": "Use both folders."}


def csv_file(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(FIELDS); writer.writerows(rows)


@pytest.fixture
def local(tmp_path):
    source, bi = tmp_path / "pending_files", tmp_path / "bi_system"
    # Identical filenames and misleading names cannot override explicit folder roles.
    changed = [list(row) for row in ROWS]; changed[0][3] = "125.00"
    csv_file(source / "same.csv", changed); csv_file(bi / "same.csv", ROWS)
    app = create_app(tmp_path / "workspace", model=Model())
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Requested-With": "FinanceFlow"}) as client:
        assert client.post("/api/setup", json={"username": "Folder Owner", "password": "123456"}).status_code == 200
        yield client, app.state.service, source, bi


def scan(local):
    client, service, source, bi = local
    response = client.post("/api/batches/scan", json={"folder": str(source), "bi_folder": str(bi)})
    assert response.status_code == 202, response.text
    batch_id = response.json()["id"]; service.future.result(timeout=60)
    return client.get(f"/api/batches/{batch_id}").json()


def test_separate_roots_fix_roles_and_reconcile_only_input_rows(local):
    client, service, source, bi = local
    batch = scan(local)
    assert batch["status"] == "scanned" and batch["input_rows"] == 2
    assert batch["bi_folder"] == str(bi)
    assert batch["source_tables"] == batch["reference_tables"] == 1
    tables = {item["input_group"]: item for item in batch["tables"]}
    assert tables["input"]["role"] == "source" and tables["bi"]["role"] == "reference"
    assert tables["input"]["reference_id"] == tables["bi"]["id"]
    assert tables["input"]["file"] == "Input/same.csv"
    assert tables["bi"]["file"] == "BI system/same.csv"
    assert client.post(f"/api/batches/{batch['id']}/run").status_code == 200
    service.future.result(timeout=60)
    done = client.get(f"/api/batches/{batch['id']}").json()
    assert done["matched_rows"] == done["archive_receipts"] == 1
    assert done["unresolved_rows"] == 1 and len(done["review_tasks"]) == 1
    book = load_workbook(io.BytesIO(client.get(f"/api/batches/{batch['id']}/report").content))
    assert any(cell.value == str(bi) for row in book["Summary"] for cell in row)


def test_bi_file_change_is_checked_against_bi_root(local):
    client, service, _, bi = local
    batch = scan(local)
    with (bi / "same.csv").open("a") as stream: stream.write("\n")
    client.post(f"/api/batches/{batch['id']}/run"); service.future.result(timeout=60)
    done = service.batches.describe(batch["id"])
    assert done["archive_receipts"] == 0
    assert any("changed after the scan" in message for issue in done["attention"] for message in issue["messages"])


def test_selected_folder_roles_cannot_be_swapped(local):
    client, _, _, _ = local
    batch = scan(local)
    entry = next(table for table in batch["tables"] if table["input_group"] == "bi")
    assert client.post(f"/api/batches/{batch['id']}/tables/{entry['id']}", json={"role": "source"}).status_code == 422


@pytest.mark.parametrize("kind", ["same", "child", "parent", "missing"])
def test_invalid_or_overlapping_folder_choices_are_rejected_before_capture(local, kind):
    client, service, source, bi = local
    child = source / "nested"; child.mkdir()
    chosen = {"same": source, "child": child, "parent": source.parent, "missing": source.parent / "missing"}[kind]
    body = {"folder": str(source), "bi_folder": str(chosen)}
    assert client.post("/api/batches/scan", json=body).status_code == 422
    assert client.post("/api/assistant/chats", json=body).status_code == 422
    assert not service.batches.list() and not service.assistant.list()


def test_empty_bi_folder_does_not_create_usable_partial_intake(local, tmp_path):
    client, service, source, _ = local
    empty = tmp_path / "empty"; empty.mkdir()
    body = {"folder": str(source), "bi_folder": str(empty)}
    assert client.post("/api/assistant/chats", json=body).status_code == 422
    response = client.post("/api/batches/scan", json=body)
    assert response.status_code == 202
    service.future.result(timeout=60)
    assert not service.batches.describe(response.json()["id"])["scan_complete"]


def test_assistant_catalog_keeps_same_names_in_distinct_groups_and_blocks_append(local):
    client, service, source, bi = local
    chat = client.post("/api/assistant/chats", json={"folder": str(source), "bi_folder": str(bi)}).json()
    assert chat["bi_folder"] == str(bi)
    assert [file["group"] for file in chat["files"]] == ["input", "bi"]
    state = json.loads(service.assistant.chat(chat["id"])["state_json"])
    with service.store.connect() as db:
        turn_id = db.execute("INSERT INTO data_turns(chat_id,instruction,status,state_json,created_at,updated_at) VALUES(?,?,?,'{}','','')", (chat["id"], "Compare folders", "completed")).lastrowid
    tools = DataTools(service.assistant, chat["id"], turn_id, state)
    tools.call("inspect_table", {"file_id": "f1"}); tools.call("inspect_table", {"file_id": "f2"})
    assert [file["group"] for file in tools.list_files()["files"]] == ["input", "bi"]
    with pytest.raises(ValueError, match="must not be appended"):
        tools.call("concatenate", {"frame_ids": ["t1", "t2"]})
    joined = tools.call("join_tables", {"left_id": "t1", "right_id": "t2", "left_on": ["supplier_id", "order_id"], "right_on": ["supplier_id", "order_id"]})
    assert joined["rows"] == 2 and state["frames"][joined["frame_id"]]["source_groups"] == ["bi", "input"]
    artifact = tools.call("export_report", {"frame_ids": [joined["frame_id"]]})
    book = load_workbook(io.BytesIO(client.get(artifact["url"]).content))
    assert [book["Source evidence"].cell(row, 1).value for row in (2, 3)] == ["Input", "BI system"]
    assert book["Summary"]["B7"].value == str(bi)
    csv_file(bi / "same.csv", [])
    assert tools.frame("t2").shape[0] == 2


def test_reconciliation_command_preserves_both_roots_and_auto_runs(local):
    client, service, source, bi = local
    response = client.post("/api/commands/prepare", json={"instruction": "Reconcile input files against BI and export a report.", "folder": str(source), "bi_folder": str(bi)})
    assert response.status_code == 200, response.text
    command = response.json(); assert command["bi_folder"] == str(bi)
    response = client.post(f"/api/commands/{command['id']}/execute")
    service.future.result(timeout=60)
    batch = service.batches.describe(response.json()["id"])
    assert batch["bi_folder"] == str(bi) and batch["input_rows"] == 2 and batch["archive_receipts"] == 1


def test_backup_restore_keeps_two_roots_and_snapshot_groups(local, tmp_path):
    client, service, source, bi = local
    batch = scan(local)
    chat = client.post("/api/assistant/chats", json={"folder": str(source), "bi_folder": str(bi)}).json()
    restored = tmp_path / "restored"; restore_backup(service.backup(), restored)
    app = create_app(restored)
    try:
        assert app.state.service.batches.describe(batch["id"])["bi_folder"] == str(bi)
        saved = app.state.service.assistant.describe(chat["id"])
        assert saved["bi_folder"] == str(bi) and saved["files"][1]["group"] == "bi"
    finally: app.state.service.executor.shutdown(wait=True)


def test_existing_workspace_migration_preserves_account_and_records(tmp_path):
    directory = tmp_path / "old"; store = Store(directory)
    with store.connect() as db:
        db.execute("INSERT INTO users VALUES(1,'Existing Owner','existing-hash','2026-10-01')")
        for table in ("batches", "commands", "data_chats"): db.execute(f"ALTER TABLE {table} DROP COLUMN bi_folder")
        db.execute("UPDATE settings SET value='3' WHERE key='schema_version'")
    migrated = Store(directory)
    assert migrated.one("SELECT username,password_hash FROM users")["password_hash"] == "existing-hash"
    assert migrated.one("SELECT value FROM settings WHERE key='schema_version'")["value"] == "4"
    for table in ("batches", "commands", "data_chats"):
        with migrated.connect() as db: assert "bi_folder" in {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
