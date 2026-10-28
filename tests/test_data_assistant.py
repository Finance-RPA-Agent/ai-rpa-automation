"""Data tools, persistent multi-step conversations and immutable evidence."""
import csv
import io
import json
from pathlib import Path
import threading

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from product.api import create_app
from product.data_tools import DataTools, number
from start_product import restore_backup

HEADERS = ["Order ID", "Supplier", "Amount", "Date", "Payment", "Currency"]
ROWS = [["001", "Acme", "100.10", "2026-01-01", "pending", "GBP"],
        ["002", "Cedar", "200.20", "2026-01-02", "paid", "GBP"],
        ["003", "Acme", "50.05", "2026-02-01", "pending", "GBP"]]


def action(tool, **arguments):
    return {"action": "tool", "tool": tool, "arguments": arguments, "message": ""}


def final(message="Done. The computed results are available."):
    return {"action": "final", "tool": None, "arguments": {}, "message": message}


class Model:
    def __init__(self): self.actions, self.contexts, self.requirements = [], [], None
    def data_plan(self, context, schema):
        from product.data_tools import SPECS
        required = self.requirements if self.requirements is not None else list(dict.fromkeys(item["tool"] for item in self.actions if item["action"] == "tool" and item["tool"] in SPECS))
        return {"supported": True, "required_tools": required, "message": "Use the requested tools."}
    def data_action(self, context, schema):
        self.contexts.append(context)
        return self.actions.pop(0) if self.actions else {"action": "ask", "message": "Which columns should I use?"}


@pytest.fixture
def local(tmp_path):
    inbox = tmp_path / "inbox"; inbox.mkdir()
    with (inbox / "orders.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream); writer.writerow(HEADERS); writer.writerows(ROWS)
    book = Workbook(); sheet = book.active; sheet.title = "Customers"
    sheet.append(["Supplier", "Region"]); sheet.append(["Acme", "North"]); sheet.append(["Cedar", "South"])
    sheet = book.create_sheet("Other"); sheet.append(["ID", "Name"]); sheet.append(["007", "Example"])
    book.save(inbox / "vendors.xlsx"); book.close()
    model = Model(); app = create_app(tmp_path / "workspace", model=model)
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Requested-With": "FinanceFlow"}) as client:
        assert client.post("/api/setup", json={"username": "Data Owner", "password": "123456"}).status_code == 200
        chat = client.post("/api/assistant/chats", json={"folder": str(inbox)}).json()
        yield client, app.state.service, model, chat, inbox


def turn(local, actions, instruction="Analyze the orders."):
    client, service, model, chat, _ = local
    model.actions = list(actions)
    response = client.post(f"/api/assistant/chats/{chat['id']}/turns", json={"instruction": instruction})
    assert response.status_code == 202, response.text
    service.future.result(timeout=60)
    return client.get(f"/api/assistant/chats/{chat['id']}").json()


def tools(local):
    result = turn(local, [])
    _, service, _, chat, _ = local
    state = json.loads(service.assistant.chat(chat["id"])["state_json"])
    return DataTools(service.assistant, chat["id"], result["turns"][-1]["id"], state)


def test_multi_step_query_follow_up_chart_and_excel(local):
    result = turn(local, [action("inspect_table", file_id="f1"),
                          action("filter_rows", frame_id="t1", conditions=[{"column": "Payment", "op": "eq", "value": "pending"}]),
                          action("aggregate", frame_id="t2", group_by=["Supplier"], metrics=[{"column": "Amount", "op": "sum", "output": "Total"}]),
                          action("create_chart", frame_id="t3", x="Supplier", y="Total", kind="bar"),
                          action("export_report", frame_ids=["t3"], title="Unpaid orders"), final()], "Summarize unpaid orders by supplier and export a chart report.")
    assert result["turns"][-1]["status"] == "completed"
    assert result["preview"]["preview"] == [["Acme", "150.15"]]
    assert result["charts"][0]["values"] == ["150.15"]
    client, service, model, chat, _ = local
    artifact = result["artifacts"][0]
    response = client.get(artifact["url"])
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book["Result 1"]["B2"].value == "150.15"
    assert len(book["Chart 1"]._charts) == 1
    assert {"Summary", "Operations", "Source evidence"}.issubset(book.sheetnames)
    followup = turn(local, [action("sort_rows", frame_id="t3", by=["Total"], descending=True, kinds={"Total": "number"}), final()], "Sort that result by total descending.")
    assert followup["turns"][-1]["status"] == "completed"
    assert model.contexts[-2]["latest_frame"] == "t3"
    assert model.contexts[-2]["previous_requests"][0]["request"].startswith("Summarize unpaid")


def test_csv_excel_multiple_sheets_and_snapshot_immutability(local):
    engine = tools(local); _, _, _, _, inbox = local
    first = engine.call("inspect_table", {"file_id": "f1"})
    assert first["preview"][0][0] == "001"
    (inbox / "orders.csv").write_text("changed", encoding="utf-8")
    assert engine.frame("t1").iloc[0, 0] == "001"
    assert engine.call("inspect_table", {"file_id": "f1"})["frame_id"] == "t1"
    excel = engine.call("inspect_table", {"file_id": "f2", "sheet": "Other"})
    assert excel["preview"] == [["007", "Example"]]


def test_clean_append_deduplicate_and_data_quality(local):
    engine = tools(local)
    first = engine.save(pd.DataFrame([["001", " Acme ", ""], ["001", "Acme", "5"], ["002", "CEDAR", "7"]], columns=["ID", "Name", "Amount"]), "Test input")
    trimmed = engine.call("normalize_text", {"frame_id": "t1", "columns": ["Name"], "mode": "trim"})
    filled = engine.call("fill_missing", {"frame_id": trimmed["frame_id"], "values": {"Amount": "0"}})
    assert engine.call("profile_table", {"frame_id": first["frame_id"]})["column_profile"][2]["missing"] == 1
    duplicates = engine.call("deduplicate", {"frame_id": filled["frame_id"], "keys": ["ID"], "duplicates_only": True})
    assert duplicates["rows"] == 2
    unique = engine.call("deduplicate", {"frame_id": filled["frame_id"], "keys": ["ID"]})
    selected = engine.call("select_columns", {"frame_id": unique["frame_id"], "columns": ["ID", "Amount"]})
    appended = engine.call("concatenate", {"frame_ids": [selected["frame_id"], first["frame_id"]], "union_columns": True})
    assert appended["rows"] == 5 and appended["columns"] == ["ID", "Amount", "Name"]


def test_join_relationship_and_missing_keys(local):
    engine = tools(local)
    engine.call("inspect_table", {"file_id": "f1"}); engine.call("inspect_table", {"file_id": "f2"})
    joined = engine.call("join_tables", {"left_id": "t1", "right_id": "t2", "left_on": ["Supplier"], "right_on": ["Supplier"], "relationship": "many_to_one"})
    assert joined["rows"] == 3 and joined["preview"][0][-1] == "North"
    with pytest.raises(ValueError, match="Duplicate join keys"):
        engine.call("join_tables", {"left_id": "t1", "right_id": "t1", "left_on": ["Supplier"], "right_on": ["Supplier"]})


def test_month_grouping_exact_arithmetic_numeric_sort_and_dates(local):
    engine = tools(local); engine.call("inspect_table", {"file_id": "f1"})
    month = engine.call("derive_column", {"frame_id": "t1", "column": "Month", "operation": "month", "left": "Date"})
    grouped = engine.call("aggregate", {"frame_id": month["frame_id"], "group_by": ["Month"], "metrics": [{"column": "Amount", "op": "sum", "output": "Total"}]})
    assert grouped["preview"] == [["2026-01", "300.30"], ["2026-02", "50.05"]]
    sorted_ = engine.call("sort_rows", {"frame_id": grouped["frame_id"], "by": ["Total"], "descending": True, "kinds": {"Total": "number"}})
    assert sorted_["preview"][0][1] == "300.30"
    calc = engine.call("derive_column", {"frame_id": "t1", "column": "Double", "operation": "multiply", "left": "Amount", "factor": "2"})
    assert calc["preview"][0][-1] == "200.20"
    filtered = engine.call("filter_rows", {"frame_id": "t1", "conditions": [{"column": "Date", "op": "gte", "value": "2026-02-01", "kind": "date"}]})
    assert filtered["rows"] == 1


def test_top_rows_and_column_renaming(local):
    engine = tools(local); engine.call("inspect_table", {"file_id": "f1"})
    sorted_ = engine.call("sort_rows", {"frame_id": "t1", "by": ["Amount"], "descending": True, "kinds": {"Amount": "number"}})
    top = engine.call("limit_rows", {"frame_id": sorted_["frame_id"], "limit": 1})
    assert top["preview"][0][0] == "002"
    renamed = engine.call("rename_columns", {"frame_id": top["frame_id"], "mapping": {"Order ID": "Order"}})
    assert renamed["columns"][0] == "Order"
    with pytest.raises(ValueError, match="unique"):
        engine.call("rename_columns", {"frame_id": "t1", "mapping": {"Order ID": "Supplier"}})


def test_invalid_numeric_groups_mixed_currency_and_ambiguous_dates(local):
    engine = tools(local)
    engine.save(pd.DataFrame([["1,2,3", "GBP", "01/02/2026"], ["2", "USD", "2026-01-02"]], columns=["Amount", "Currency", "Date"]), "Input")
    with pytest.raises(ValueError, match="Multiple currencies"):
        engine.call("aggregate", {"frame_id": "t1", "metrics": [{"column": "Amount", "op": "sum", "output": "Total"}]})
    with pytest.raises(ValueError, match="unambiguous number"):
        engine.call("aggregate", {"frame_id": "t1", "group_by": ["Currency"], "metrics": [{"column": "Amount", "op": "sum", "output": "Total"}]})
    with pytest.raises(ValueError, match="Dates must use"):
        engine.call("derive_column", {"frame_id": "t1", "column": "Month", "operation": "month", "left": "Date"})


def test_formula_text_and_precise_identifiers_remain_inert_in_reports(local):
    engine = tools(local)
    engine.save(pd.DataFrame([["00001", '=HYPERLINK("https://example.invalid")', "10000000000000.01"]], columns=["ID", "Name", "Amount"]), "Input")
    artifact = engine.call("export_report", {"frame_ids": ["t1"]})
    book = load_workbook(io.BytesIO(local[0].get(artifact["url"]).content))
    assert book["Result 1"]["B2"].data_type == "s"
    assert book["Result 1"]["A2"].value == "00001"
    assert book["Result 1"]["C2"].value == "10000000000000.01"


def test_unknown_tools_arguments_ids_and_integrity(local):
    engine = tools(local)
    for tool, args in [("python", {"code": "anything"}), ("inspect_table", {"file_id": "../../secret"}), ("inspect_table", {"file_id": "f1", "path": "C:/"}), ("select_columns", {"frame_id": "missing", "columns": ["ID"]})]:
        with pytest.raises(ValueError): engine.call(tool, args)
    engine.call("inspect_table", {"file_id": "f1"})
    engine.service.path(engine.state["frames"]["t1"]["path"]).write_text("changed")
    with pytest.raises(ValueError, match="integrity"): engine.frame("t1")


def test_model_tool_error_can_recover_and_unsupported_action_cannot_run(local):
    result = turn(local, [action("shell", command="anything"), action("inspect_table", file_id="f1"), final()])
    assert result["turns"][-1]["status"] == "completed"
    assert result["turns"][-1]["steps"][0]["status"] == "error"
    assert result["preview"]["rows"] == 3


def test_step_budget_and_error_budget_keep_results(local):
    result = turn(local, [action("inspect_table", file_id="f1")] * 12)
    assert result["turns"][-1]["status"] == "failed" and "12-step" in result["turns"][-1]["error"]
    assert result["preview"]["rows"] == 3
    result = turn(local, [action("unknown")] * 3)
    assert result["turns"][-1]["status"] == "failed" and "Three tool" in result["turns"][-1]["error"]


def test_model_question_is_not_reported_as_completed(local):
    result = turn(local, [action("list_files"), final("Which date column should I use?")])
    assert result["turns"][-1]["status"] == "needs_input"


def test_claimed_report_without_real_report_is_rejected(local):
    local[2].requirements = ["inspect_table", "aggregate", "export_report"]
    result = turn(local, [action("inspect_table", file_id="f1"), final("The grouped report has been exported.")])
    assert result["turns"][-1]["status"] == "needs_input"
    assert not result["artifacts"]
    assert result["turns"][-1]["steps"][-1]["tool"] == "completion_check"


def test_unbacked_currency_symbols_are_not_added_to_results(local):
    result = turn(local, [action("inspect_table", file_id="f1"), final("The value is $100.10; the source currency is £100.10.")])
    message = result["turns"][-1]["message"]
    assert "$" not in message and "£100.10" in message


def test_explicit_first_row_cannot_be_omitted_by_model_plan(local):
    local[2].requirements = ["inspect_table", "export_report"]
    result = turn(local, [action("inspect_table", file_id="f1"), action("export_report", frame_ids=["t1"]), final("The first row was saved.")], "只保留第一行，另存为 Excel。")
    assert result["turns"][-1]["status"] == "needs_input"
    assert "limit_rows" in result["turns"][-1]["required_tools"]
    assert result["artifacts"][0]["row_count"] == 3


def test_currency_evidence_survives_column_selection(local):
    engine = tools(local)
    engine.save(pd.DataFrame([["1", "GBP"], ["2", "USD"]], columns=["Amount", "Currency"]), "Mixed currency input")
    selected = engine.call("select_columns", {"frame_id": "t1", "columns": ["Amount"]})
    with pytest.raises(ValueError, match="source evidence"):
        engine.call("aggregate", {"frame_id": selected["frame_id"], "metrics": [{"column": "Amount", "op": "sum", "output": "Total"}]})


def test_sort_must_precede_top_row_selection(local):
    local[2].requirements = ["inspect_table", "sort_rows", "limit_rows"]
    result = turn(local, [action("inspect_table", file_id="f1"), action("limit_rows", frame_id="t1", limit=1),
        action("sort_rows", frame_id="t1", by=["Amount"], descending=True, kinds={"Amount": "number"}),
        action("limit_rows", frame_id="t2", limit=1), final()], "Show the top 1 orders by amount, descending.")
    assert result["turns"][-1]["status"] == "completed"
    assert result["turns"][-1]["steps"][1]["status"] == "error"
    assert result["preview"]["preview"][0][0] == "002"


def test_chart_report_must_actually_contain_requested_chart(local):
    result = turn(local, [action("inspect_table", file_id="f1"),
        action("aggregate", frame_id="t1", group_by=["Supplier"], metrics=[{"column": "Amount", "op": "sum", "output": "Total"}]),
        action("create_chart", frame_id="t2", x="Supplier", y="Total"),
        action("export_report", frame_ids=["t1"]), final("The report contains the chart.")])
    assert result["turns"][-1]["status"] == "needs_input"
    assert not result["artifacts"][0]["chart_ids"]
    assert result["turns"][-1]["steps"][-1]["tool"] == "completion_check"


def test_cancel_worker_exclusion_and_backup(local, monkeypatch):
    client, service, model, chat, inbox = local
    entered, release = threading.Event(), threading.Event()
    def blocked(context, schema):
        entered.set(); assert release.wait(10); return action("inspect_table", file_id="f1")
    model.data_action = blocked
    response = client.post(f"/api/assistant/chats/{chat['id']}/turns", json={"instruction": "Read the table."})
    turn_id = response.json()["id"]; assert entered.wait(10)
    try:
        assert client.post("/api/batches/scan", json={"folder": str(inbox)}).status_code == 409
        assert client.post("/api/backup").status_code == 409
        assert client.post(f"/api/assistant/chats/{chat['id']}/turns", json={"instruction": "Another request."}).status_code == 409
        assert client.post(f"/api/assistant/turns/{turn_id}/cancel").status_code == 200
    finally: release.set()
    service.future.result(timeout=60)
    result = client.get(f"/api/assistant/chats/{chat['id']}").json()
    assert result["turns"][-1]["status"] == "cancelled" and not result["frames"]
    assert client.post("/api/backup").status_code == 200


def test_chat_and_artifact_survive_backup_restore(local, tmp_path):
    result = turn(local, [action("inspect_table", file_id="f1"), action("export_report", frame_ids=["t1"]), final()])
    service = local[1]; target = tmp_path / "restored"
    restore_backup(service.backup(), target)
    restored = create_app(target)
    try:
        chat = restored.state.service.assistant.describe(local[3]["id"])
        assert chat["preview"]["preview"][0][0] == "001"
        assert restored.state.service.assistant.artifact(result["artifacts"][0]["id"]).is_file()
    finally: restored.state.service.executor.shutdown(wait=True)


def test_authentication_csrf_and_path_boundaries(local):
    client, service, _, chat, _ = local
    assert client.post("/api/assistant/chats", json={"folder": str(service.store.root)}).status_code == 422
    assert client.post("/api/assistant/chats", json={"folder": "../relative"}).status_code == 422
    assert client.post(f"/api/assistant/chats/{chat['id']}/turns", json={"instruction": "   "}).status_code == 422
    assert client.post(f"/api/assistant/chats/{chat['id']}/turns", headers={"X-Requested-With": "other"}, json={"instruction": "Read files"}).status_code == 403
    client.post("/api/logout")
    for path in ("/api/assistant/chats", f"/api/assistant/chats/{chat['id']}", "/api/assistant/artifacts/anything"):
        assert client.get(path).status_code == 401


def test_report_integrity_and_native_export_fallback(local, monkeypatch):
    result = turn(local, [action("inspect_table", file_id="f1"), action("export_report", frame_ids=["t1"]), final()])
    import product.native_reports
    monkeypatch.setattr(product.native_reports, "save_report", lambda *args: {"handled": False})
    client, service, _, _, _ = local
    artifact = result["artifacts"][0]
    saved = client.post(artifact["url"] + "/export")
    assert saved.json()["download_url"] == artifact["url"]
    service.assistant.artifact(artifact["id"]).write_bytes(b"changed")
    assert client.get(artifact["url"]).status_code == 409


@pytest.mark.parametrize("value", ["1,2,3", "NaN", "Infinity", "$12", "1e300", "", "1.234,56"])
def test_ambiguous_or_nonfinite_numbers_are_rejected(value):
    with pytest.raises(ValueError): number(value)
