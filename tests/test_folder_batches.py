"""End-to-end folder intake, safe partial reconciliation and Excel audit reports."""
import csv
import io
import json
from pathlib import Path
import threading

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from product.api import create_app
from product.service import FIELDS
from start_product import restore_backup

ROWS = [
    ["001", "Acme", "0001", "100.00", "2026-10-01", "pending", "active"],
    ["002", "Cedar", "0002", "200.00", "2026-10-02", "paid", "active"],
    ["003", "Delta", "0003", "300.00", "2026-10-03", "pending", "active"],
]
ALIASES = ["Vendor ID", "Vendor Name", "PO Number", "Total Amount", "Invoice Date", "Payment State", "Record State"]


def csv_file(path, rows, headers=FIELDS):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


@pytest.fixture
def local(tmp_path):
    app = create_app(tmp_path / "workspace")
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Requested-With": "FinanceFlow"}) as client:
        assert client.post("/api/setup", json={"username": "Batch Owner", "password": "123456"}).status_code == 200
        yield client, app.state.service, tmp_path / "inbox"


def scan(local):
    client, service, inbox = local
    response = client.post("/api/batches/scan", json={"folder": str(inbox)})
    assert response.status_code == 202, response.text
    batch_id = response.json()["id"]
    service.future.result(timeout=60)
    view = client.get(f"/api/batches/{batch_id}").json()
    assert view["status"] == "scanned", view
    return view


def run(local, batch_id):
    client, service, _ = local
    response = client.post(f"/api/batches/{batch_id}/run")
    assert response.status_code == 200, response.text
    service.future.result(timeout=60)
    return client.get(f"/api/batches/{batch_id}").json()


def report(client, batch_id):
    response = client.get(f"/api/batches/{batch_id}/report")
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content), data_only=False)


def test_normal_rows_archive_automatically_only_exceptions_wait(local):
    client, service, inbox = local
    source = [list(row) for row in ROWS]
    source[0][3] = "125.00"
    source[1][1] = ""
    csv_file(inbox / "source" / "orders.csv", source, ALIASES)
    csv_file(inbox / "reference" / "bi.csv", ROWS)
    batch = scan(local)
    assert batch["runnable_tables"] == 1
    done = run(local, batch["id"])
    assert done["status"] == "needs_review", done
    assert done["matched_rows"] == done["archive_receipts"] == 1
    assert done["unresolved_rows"] == 2
    assert len(done["review_tasks"]) == 1
    job_id = done["review_tasks"][0]["job_id"]
    task = client.get(f"/api/tasks/{job_id}").json()
    assert task["issues"][0]["original_row"] == 3
    workbook = report(client, batch["id"])
    assert workbook.sheetnames == ["Summary", "Exceptions", "Reconciliation", "File log", "Review history"]
    findings = [row[8] for row in workbook["Exceptions"].iter_rows(min_row=6, values_only=True)]
    assert any("amount" in str(value).lower() for value in findings)
    assert any("supplier_name" in str(value) for value in findings)
    workbook.close()
    # Human review remains specific and revision-bound; corrections are rechecked.
    for _ in range(3):
        task = client.get(f"/api/tasks/{job_id}").json()
        if not task["issues"]:
            break
        response = client.post(f"/api/tasks/{job_id}/review", json={"revision": task["revision"], "evidence_digest": task["evidence_digest"],
            "changes": [{"issue_id": issue["id"], "mode": "reference", "value": None} for issue in task["issues"]],
            "reason": "Verified each original order against the BI export."})
        assert response.status_code == 200, response.text
    assert client.post(f"/api/tasks/{job_id}/run").status_code == 200
    service.future.result(timeout=60)
    final = client.get(f"/api/batches/{batch['id']}").json()
    assert final["status"] == "completed" and final["matched_rows"] == final["archive_receipts"] == 3
    workbook = report(client, batch["id"])
    assert workbook["Summary"]["B4"].value == "FINAL - completed"
    assert workbook["Exceptions"]["A6"].value == "No outstanding items."
    assert workbook["Review history"].max_row == 7
    workbook.close()


def test_excel_sheets_header_detection_text_ids_and_formula_safety(local):
    client, _, inbox = local
    inbox.mkdir()
    workbook = Workbook()
    workbook.active.title = "Orders"
    workbook.active.append(["October orders"])
    workbook.active.append(ALIASES)
    rows = [list(row) for row in ROWS]
    rows[0][1] = '=HYPERLINK("https://example.invalid","untrusted")'
    rows[2][3] = "10000000000000.01"
    for row in rows:
        workbook.active.append(row)
        workbook.active.cell(workbook.active.max_row, 2).data_type = "s"
    reference = workbook.create_sheet("BI")
    reference.append(FIELDS)
    for row in rows:
        reference.append(row)
        reference.cell(reference.max_row, 2).data_type = "s"
    workbook.save(inbox / "monthly.xlsx")
    workbook.close()
    batch = scan(local)
    assert batch["source_tables"] == batch["reference_tables"] == 1
    done = run(local, batch["id"])
    assert done["status"] == "completed", done
    output = report(client, batch["id"])
    sheet = output["Reconciliation"]
    assert sheet["D6"].value == 3
    assert sheet["E6"].value == "001" and sheet["F6"].value == "0001"
    assert sheet["Q6"].value.startswith("=HYPERLINK") and sheet["Q6"].data_type == "s"
    assert sheet["G8"].value == "10000000000000.01" and sheet["G8"].data_type == "s"
    assert not any(cell.data_type == "f" for sheet in output for row in sheet for cell in row)
    output.close()


def test_ambiguous_reference_requires_explicit_choice(local):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference_a.csv", ROWS)
    csv_file(inbox / "reference_b.csv", ROWS)
    batch = scan(local)
    source = next(t for t in batch["tables"] if t["role"] == "source")
    assert source["reference_id"] is None and batch["runnable_tables"] == 0
    assert client.post(f"/api/batches/{batch['id']}/run").status_code == 409
    reference = next(t for t in batch["tables"] if t["role"] == "reference")
    response = client.post(f"/api/batches/{batch['id']}/tables/{source['id']}", json={"role": "source", "reference_id": reference["id"]})
    assert response.status_code == 200, response.text
    assert run(local, batch["id"])["status"] == "completed"


def test_changed_reference_and_cross_table_duplicates_never_auto_archive(local):
    client, _, inbox = local
    csv_file(inbox / "orders_a.csv", ROWS[:2])
    csv_file(inbox / "orders_b.csv", [ROWS[0], ROWS[2]])
    csv_file(inbox / "reference.csv", ROWS)
    batch = scan(local)
    done = run(local, batch["id"])
    assert done["archive_receipts"] == 2 and done["matched_rows"] == 2
    assert done["status"] == "needs_review" and len(done["attention"]) == 2
    workbook = report(client, batch["id"])
    results = [row[0] for row in workbook["Reconciliation"].iter_rows(min_row=6, values_only=True)]
    assert results.count("Source correction required") == 2
    workbook.close()
    another = scan(local)
    csv_file(inbox / "reference.csv", [ROWS[0]])
    stopped = run(local, another["id"])
    assert stopped["matched_rows"] == stopped["archive_receipts"] == 0
    assert any("changed" in str(item["messages"]) for item in stopped["attention"])


def test_identical_rescan_reuses_receipts_and_backup_restores_batch(local, tmp_path):
    client, service, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    original = run(local, scan(local)["id"])
    repeated = run(local, scan(local)["id"])
    assert original["archive_receipts"] == repeated["archive_receipts"] == 3
    assert repeated["reused_tasks"] == 1
    assert len(client.get("/api/archive").json()["records"]) == 3
    archive = tmp_path / "batch-backup.zip"
    archive.write_bytes(client.post("/api/backup").content)
    restored_dir = tmp_path / "restored"
    restore_backup(archive, restored_dir)
    app = create_app(restored_dir)
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Requested-With": "FinanceFlow"}) as restored:
        assert restored.post("/api/login", json={"username": "Batch Owner", "password": "123456"}).status_code == 200
        assert restored.get(f"/api/batches/{original['id']}").json()["status"] == "completed"
        workbook = report(restored, original["id"])
        assert workbook["Reconciliation"].max_row == 8
        workbook.close()


def test_unknown_columns_can_be_mapped_and_broken_files_ignored(local):
    client, _, inbox = local
    names = ["Code", "Company", "Document", "Value", "Day", "Settlement", "State"]
    csv_file(inbox / "input.csv", ROWS, names)
    csv_file(inbox / "reference.csv", ROWS)
    (inbox / "broken.xlsx").write_bytes(b"not a workbook")
    batch = scan(local)
    source = next(t for t in batch["tables"] if t["file"] == "input.csv")
    response = client.post(f"/api/batches/{batch['id']}/tables/{source['id']}", json={"role": "source", "mapping": dict(zip(FIELDS, names))})
    assert response.status_code == 200, response.text
    broken = next(t for t in batch["tables"] if t["file"] == "broken.xlsx")
    assert client.post(f"/api/batches/{batch['id']}/tables/{broken['id']}", json={"role": "skip"}).status_code == 200
    assert run(local, batch["id"])["status"] == "completed"


def test_currency_conflicts_are_explicit_and_auth_is_required(local):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", [row + ["GBP"] for row in ROWS], FIELDS + ["Currency"])
    csv_file(inbox / "reference.csv", [row + ["USD"] for row in ROWS], FIELDS + ["Currency"])
    batch = scan(local)
    assert batch["runnable_tables"] == 0
    assert any("currencies differ" in str(item["messages"]) for item in batch["attention"])
    client.post("/api/logout")
    assert client.get("/api/batches").status_code == 401
    assert client.get(f"/api/batches/{batch['id']}/report").status_code == 401


def test_active_scan_blocks_other_mutations_and_cancels_safely(local, monkeypatch):
    client, service, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    started, release = threading.Event(), threading.Event()
    original = service.batches._scan
    def slow(batch_id):
        started.set()
        release.wait(10)
        return original(batch_id)
    monkeypatch.setattr(service.batches, "_scan", slow)
    response = client.post("/api/batches/scan", json={"folder": str(inbox)})
    assert response.status_code == 202
    batch_id = response.json()["id"]
    try:
        assert started.wait(5)
        assert client.post("/api/backup").status_code == 409
        assert client.post("/api/batches/scan", json={"folder": str(inbox)}).status_code == 409
        assert client.post(f"/api/batches/{batch_id}/cancel").status_code == 200
    finally:
        release.set()
    service.future.result(timeout=20)
    assert client.get(f"/api/batches/{batch_id}").json()["status"] == "cancelled"
    assert client.get("/api/archive").json()["records"] == []


def test_incomplete_scan_cannot_be_executed_or_reconfigured(local, monkeypatch):
    client, service, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    original = service.batches._scan
    def interrupted(batch_id):
        original(batch_id)
        raise RuntimeError("The scan was interrupted before completion.")
    monkeypatch.setattr(service.batches, "_scan", interrupted)
    response = client.post("/api/batches/scan", json={"folder": str(inbox)})
    batch_id = response.json()["id"]
    service.future.result(timeout=30)
    batch = client.get(f"/api/batches/{batch_id}").json()
    assert batch["status"] == "failed" and not batch["scan_complete"]
    assert batch["runnable_tables"] == 0
    assert client.post(f"/api/batches/{batch_id}/run").status_code == 409
    entry = batch["tables"][0]
    assert client.post(f"/api/batches/{batch_id}/tables/{entry['id']}", json={"role": "skip"}).status_code == 409
    assert client.get("/api/archive").json()["records"] == []


def test_used_reference_and_pairing_remain_immutable(local):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    csv_file(inbox / "unclassified.csv", ROWS)
    batch = scan(local)
    source = next(t for t in batch["tables"] if t["role"] == "source")
    reference_id = source["reference_id"]
    run(local, batch["id"])
    assert client.post(f"/api/batches/{batch['id']}/tables/{reference_id}", json={"role": "skip"}).status_code == 409
    unknown = next(t for t in batch["tables"] if t["role"] == "unknown")
    response = client.post(f"/api/batches/{batch['id']}/tables/{unknown['id']}", json={"role": "reference"})
    assert response.status_code == 200, response.text
    source_after = next(t for t in response.json()["tables"] if t["id"] == source["id"])
    assert source_after["reference_id"] == reference_id
    assert response.json()["status"] == "completed"
    used_reference = next(t for t in response.json()["tables"] if t["id"] == reference_id)
    assert used_reference["has_jobs"]


def test_reused_task_preserves_current_batch_source_row_context(local):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    first = scan(local)
    run(local, first["id"])
    original = (inbox / "orders.csv").read_text(encoding="utf-8-sig")
    (inbox / "orders.csv").write_text("October orders\n" + original, encoding="utf-8-sig")
    second = run(local, scan(local)["id"])
    assert second["reused_tasks"] == 1
    job_id = client.get("/api/tasks").json()["tasks"][0]["id"]
    old = client.get(f"/api/tasks/{job_id}?batch_id={first['id']}").json()
    current = client.get(f"/api/tasks/{job_id}?batch_id={second['id']}").json()
    assert old["records"][0]["original_row"] == 2
    assert current["records"][0]["original_row"] == 3 and current["batch_id"] == second["id"]
    assert client.get(f"/api/tasks/{job_id}?batch_id=9999").status_code == 404


def test_reference_currency_is_part_of_reuse_identity(local):
    _, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", [row + ["GBP"] for row in ROWS], FIELDS + ["Currency"])
    first = run(local, scan(local)["id"])
    assert first["archive_receipts"] == 3
    csv_file(inbox / "reference.csv", [row + ["USD"] for row in ROWS], FIELDS + ["Currency"])
    second = run(local, scan(local)["id"])
    assert second["reused_tasks"] == 0 and second["archive_receipts"] == 3


def test_report_preserves_invalid_amount_text_without_inventing_a_difference(local):
    client, _, inbox = local
    rows = [list(row) for row in ROWS]
    rows[0][3] = "1,2,3"
    csv_file(inbox / "orders.csv", rows)
    csv_file(inbox / "reference.csv", ROWS)
    batch = run(local, scan(local)["id"])
    assert batch["matched_rows"] == 2 and batch["unresolved_rows"] == 1
    workbook = report(client, batch["id"])
    assert workbook["Reconciliation"]["G6"].value == "1,2,3"
    assert workbook["Reconciliation"]["H6"].value == "1,2,3"
    assert workbook["Reconciliation"]["J6"].value is None
    workbook.close()


def test_prepared_browser_export_is_authenticated_and_readable(local, monkeypatch):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    batch = run(local, scan(local)["id"])
    monkeypatch.setattr("product.native_reports.webview.windows", [])
    response = client.post(f"/api/batches/{batch['id']}/export")
    assert response.status_code == 200, response.text
    assert response.json()["handled"] is False
    path = response.json()["download_url"]
    workbook_response = client.get(path)
    assert workbook_response.status_code == 200
    workbook = load_workbook(io.BytesIO(workbook_response.content))
    assert workbook["Reconciliation"].max_row == 8
    workbook.close()
    assert client.get(f"/api/batches/{batch['id']}/report-file/finance-flow.sqlite").status_code == 404
    client.post("/api/logout")
    assert client.get(path).status_code == 401


def test_native_report_save_and_cancel_do_not_change_reconciliation(local, monkeypatch, tmp_path):
    client, _, inbox = local
    csv_file(inbox / "orders.csv", ROWS)
    csv_file(inbox / "reference.csv", ROWS)
    batch = run(local, scan(local)["id"])
    target = tmp_path / "chosen-report.xlsx"
    class Window:
        selected = (str(target),)
        def create_file_dialog(self, dialog_type, **options):
            assert options["save_filename"].endswith(".xlsx")
            return self.selected
    window = Window()
    monkeypatch.setattr("product.native_reports.webview.windows", [window])
    response = client.post(f"/api/batches/{batch['id']}/export")
    assert response.status_code == 200 and response.json()["saved"]
    workbook = load_workbook(target)
    assert workbook["Summary"]["B4"].value == "FINAL - completed"
    workbook.close()
    original_bytes = target.read_bytes()
    window.selected = None
    response = client.post(f"/api/batches/{batch['id']}/export")
    assert response.status_code == 200 and response.json()["cancelled"]
    assert target.read_bytes() == original_bytes
    assert client.get(f"/api/batches/{batch['id']}").json()["archive_receipts"] == 3
